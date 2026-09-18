from pathlib import Path
from io import BytesIO

from docx import Document
from flask import Flask
from werkzeug.datastructures import FileStorage

from database import database
from modules.salud_nutricion.entregables import EntregablesSaludNutricionService
from modules.salud_nutricion.integral import SaludNutricionIntegralService
from modules.seguridad.tenant_context import tenant_context
from modules.sqlalchemy_compat import CoreCompatRepository


def test_informe_mensual_usa_solo_actividades_confirmadas(tmp_path):
    db_path = tmp_path / 'entregables.db'
    app = Flask(__name__)
    app.config.update(
        DATABASE_URL=f"sqlite:///{db_path.as_posix()}",
        DATABASE_PATH=str(db_path),
        SQLALCHEMY_ENGINE_OPTIONS={},
    )
    database.configure(app)
    repo = CoreCompatRepository()
    repo.log = lambda *args, **kwargs: None
    service = EntregablesSaludNutricionService(repo, str(tmp_path / 'salidas'), str(tmp_path / 'cargas'))
    integral = SaludNutricionIntegralService(repo, tmp_path / 'datos')
    try:
        with tenant_context(1, role='SUPERADMIN', username='prueba'):
            service.init_schema()
            integral.init_schema()
            result = service.crear_mes({'mes': 8, 'anio': 2026, 'uds': 'Kiphara', 'coordinador': 'Coordinadora Uno', 'responsable': 'Profesional SN'})
            assert result['creados'] == 16
            rows = service.listar({'mes': 8, 'anio': 2026, 'uds': 'Kiphara', 'coordinador': 'Coordinadora Uno'})['entregables']
            assert len(rows) == 16
            assert all(row['coordinador'] == 'Coordinadora Uno' for row in rows)
            assert service.listar({'mes': 8, 'anio': 2026, 'uds': 'Kiphara', 'coordinador': 'Otra Coordinadora'})['entregables'] == []
            lavado = next(row for row in rows if row['codigo'] == 'E02_LAVADO_MANOS')

            service.guardar_actividad(lavado['id'], {
                'fecha_actividad': '2026-08-20',
                'lugar': 'Kiphara',
                'responsable': 'Profesional SN',
                'objetivo': 'Promover el lavado de manos.',
                'desarrollo': 'Texto de borrador que no debe publicarse.',
                'confirmado': False,
            })
            service.guardar_actividad(lavado['id'], {
                'fecha_actividad': '2026-08-21',
                'lugar': 'Kiphara',
                'dirigido_a': 'Familias',
                'responsable': 'Profesional SN',
                'objetivo': 'Promover el lavado de manos.',
                'agenda': ['Bienvenida', 'Demostración', 'Compromisos'],
                'desarrollo': 'La actividad fue realizada con demostración práctica.',
                'resultados': 'Las familias practicaron la técnica.',
                'compromisos': 'Reforzar la práctica en el hogar.',
                'participantes_total': 18,
                'confirmado': True,
            })

            template_doc = Document()
            template_doc.add_paragraph('ENCABEZADO OFICIAL PRESERVADO')
            template_table = template_doc.add_table(rows=2, cols=2)
            template_table.cell(0, 0).text = 'Fecha:'
            template_table.cell(0, 1).text = ''
            template_table.cell(1, 0).text = 'Desarrollo de la actividad:'
            template_table.cell(1, 1).text = ''
            template_buffer = BytesIO()
            template_doc.save(template_buffer)
            template_buffer.seek(0)
            service.registrar_plantilla('E02_LAVADO_MANOS', 'acta', FileStorage(
                stream=template_buffer, filename='acta_oficial.docx'
            ))
            acta = service.generar_acta(lavado['id'])
            informe = service.generar_informe({'mes': 8, 'anio': 2026, 'uds': 'Kiphara'})
            generated_acta = Document(acta['ruta_archivo'])
            acta_text = '\n'.join(
                [p.text for p in generated_acta.paragraphs]
                + [cell.text for table in generated_acta.tables for row in table.rows for cell in row.cells]
            )
            informe_text = '\n'.join(p.text for p in Document(informe['ruta_archivo']).paragraphs)
            assert 'La actividad fue realizada con demostración práctica.' in acta_text
            assert 'Texto de borrador que no debe publicarse.' not in acta_text
            assert '1 actividades confirmadas' in informe_text
            assert '18 participaciones reportadas' in informe_text
            assert Path(acta['ruta_archivo']).exists()
            assert Path(informe['ruta_archivo']).exists()

            period = integral.save_monthly_period(1, {'anio_mes': '2026-08'}, {'id': 7, 'username': 'profesional'})
            assert 'Lavado de manos' in period['temas']
            assert period['variables']['actividades_confirmadas'] == 1
            assert period['variables']['participaciones_reportadas'] == 18
            assert period['variables']['fuente_automatica'] == 'Entregables Salud y Nutrición'

            # Una instalación anterior conserva IDs y relaciones al migrar los nombres provisionales.
            e02_id = repo.fetch_one("SELECT id FROM sn_entregables_catalogo WHERE codigo='E02_LAVADO_MANOS'")['id']
            repo.execute("UPDATE sn_entregables_catalogo SET codigo='E02_AAVN' WHERE id=?", (e02_id,))
            repo.execute("UPDATE sn_entregables_mes SET codigo='E02_AAVN' WHERE catalogo_id=?", (e02_id,))
            service.init_schema()
            migrated = repo.fetch_one("SELECT id,codigo FROM sn_entregables_catalogo WHERE codigo='E02_LAVADO_MANOS'")
            assert migrated['id'] == e02_id
            assert repo.fetch_one("SELECT COUNT(*) total FROM sn_entregables_catalogo")['total'] == 16
            assert repo.fetch_one("SELECT codigo FROM sn_entregables_mes WHERE catalogo_id=? LIMIT 1", (e02_id,))['codigo'] == 'E02_LAVADO_MANOS'
    finally:
        if database.engine is not None:
            database.engine.dispose()
