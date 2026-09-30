from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from modules.programas_icbf.generator import (
    F2_ROWS,
    RFPP_ROWS,
    chunks,
    generate_f2,
    generate_rfpp,
    partition_compatible,
    sort_participants,
)

RFPP_DEFAULT = r"C:\Users\kioskUser0\Downloads\PROGRAMA DE NUTRICION\RFPP. LISTO PARA IMPRIMIR..xlsx"
F2_DEFAULT = r"C:\Users\kioskUser0\Downloads\PROGRAMA DE NUTRICION\F2 BIENESTARINA DE MILDIAS.xlsx"


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def participant(index: int, first: str | None = None, last: str | None = None, category: str = "NIÑO") -> dict:
    return {
        "id": index,
        "canonical": {
            "participante.primer_nombre": first or f"NOMBRE{index:02d}",
            "participante.segundo_nombre": "",
            "participante.primer_apellido": last or f"APELLIDO{index:02d}",
            "participante.segundo_apellido": "",
            "participante.tipo_documento": "RC",
            "participante.numero_documento": f"DOC{index:04d}",
            "participante.grupo_etario": category,
            "entidad_contratista.nombre": "OPERADOR UNO",
            "regional.nombre": "REGIONAL",
            "municipio.nombre": "MUNICIPIO",
            "centro_zonal.nombre": "CENTRO",
            "unidad.codigo": "UDS-1",
            "unidad.nombre": "UDS UNO",
        },
    }


def structure(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, keep_links=True)
    return {
        "sheets": wb.sheetnames,
        "states": [wb[name].sheet_state for name in wb.sheetnames],
        "merges": {name: sorted(str(item) for item in wb[name].merged_cells.ranges) for name in wb.sheetnames},
        "images": {name: len(wb[name]._images) for name in wb.sheetnames},
        "orientation": {name: wb[name].page_setup.orientation for name in wb.sheetnames},
    }


def test_order_and_pagination():
    mixed = [
        participant(4, "Ángela", "Zuluaga", "GESTANTE"),
        participant(2, "Nora", "Ñustes", "NIÑO"),
        participant(3, "Nora", "Navarro", "GESTANTE"),
        participant(1, "Beatriz", "Álvarez", "NIÑA"),
    ]
    by_last = sort_participants(mixed, "PRIMER_APELLIDO")
    require([item["id"] for item in by_last] == [1, 3, 2, 4], "Orden español por apellido incorrecto")
    by_first = sort_participants(mixed, "PRIMER_NOMBRE")
    require([item["id"] for item in by_first] == [4, 1, 3, 2], "Orden por nombre o desempate estable incorrecto")
    require([item["id"] for item in sort_participants([{**item, "canonical": {**item["canonical"], "participante.grupo_etario": "OTRO"}} for item in mixed], "PRIMER_APELLIDO")] == [1, 3, 2, 4], "La categoría alteró el orden mixto")
    blocked = False
    try:
        sort_participants(mixed, None)
    except ValueError:
        blocked = True
    require(blocked, "Generó una secuencia definitiva con criterio pendiente")
    for total, rfpp_pages, f2_pages in ((1, 1, 1), (14, 1, 1), (15, 1, 1), (20, 1, 1), (21, 2, 1), (30, 2, 1), (31, 2, 2)):
        items = [participant(i) for i in range(1, total + 1)]
        require(len(chunks(items, 20)) == rfpp_pages, f"Paginación RFPP incorrecta para {total}")
        require(len(chunks(items, 30)) == f2_pages, f"Paginación F2 incorrecta para {total}")
    blocked = False
    try:
        chunks([], 20)
    except ValueError:
        blocked = True
    require(blocked, "Cero participantes produjo un éxito ficticio")
    split = mixed + [participant(9)]
    split[-1]["canonical"]["unidad.codigo"] = "UDS-2"
    require(sorted(len(group) for group in partition_compatible(split)) == [1, 4], "Mezcló encabezados incompatibles")


def test_real_templates():
    rfpp = Path(os.getenv("ICBF_RFPP_TEMPLATE_PATH", RFPP_DEFAULT))
    f2 = Path(os.getenv("ICBF_F2_TEMPLATE_PATH", F2_DEFAULT))
    if not rfpp.is_file() or not f2.is_file():
        raise RuntimeError("BLOQUEADO: no están disponibles las plantillas RFPP/F2 para la prueba de preservación")
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        rfpp_out = root / "rfpp.xlsx"
        f2_out = root / "f2.xlsx"
        rfpp_people = [participant(i) for i in range(1, 21)]
        f2_people = [participant(i) for i in range(1, 31)]
        generate_rfpp(rfpp, rfpp_out, rfpp_people, {"regional": "REGIONAL", "centro_zonal": "CENTRO", "municipio": "MUNICIPIO", "operador": "OPERADOR", "unidad_nombre": "UDS UNO"})
        confirmed_fact = {1: {"fecha_entrega": "2026-09-30", "receptor_nombre": "RECEPTOR", "receptor_documento": "DOC-R", "receptor_parentesco": "MADRE", "items": [{"producto": "Bienestarina Más", "lote": "LOTE-1", "unidades": 2}]}}
        generate_f2(f2, f2_out, f2_people, {"regional": "REGIONAL", "centro_zonal": "CENTRO", "municipio": "MUNICIPIO", "punto_nombre": "UDS UNO"}, confirmed_fact)
        require(structure(rfpp) == structure(rfpp_out), "La copia RFPP cambió hojas, visibilidad, imágenes, combinaciones u orientación")
        require(structure(f2) == structure(f2_out), "La copia F2 cambió hojas, visibilidad, imágenes, combinaciones u orientación")
        rfpp_wb = load_workbook(rfpp_out, data_only=False, keep_links=True)
        rfpp_ws = rfpp_wb["Formato_asist y entrega RFPP"]
        require(rfpp_ws["B14"].value == "NOMBRE01 APELLIDO01" and rfpp_ws["B33"].value == "NOMBRE20 APELLIDO20", "RFPP perdió continuidad 1–20")
        require(rfpp_ws["G14"].value == "DOC0001" and rfpp_ws["G14"].number_format == "@", "RFPP no conservó documento como texto")
        for row in RFPP_ROWS:
            require(all(rfpp_ws.cell(row, column).value in (None, "") for column in (10, 11, 12, 13, 14, 15, 16, 17, 18, 22)), "RFPP inventó asistencia, recepción, fecha o firma")
        f2_wb = load_workbook(f2_out, data_only=False, keep_links=True)
        f2_ws = f2_wb["Formato_Ajustado_Oficio"]
        require(f2_ws["B11"].value == "NOMBRE01" and f2_ws["B24"].value == "NOMBRE14", "F2 perdió el bloque 1–14")
        require(f2_ws["B32"].value == "NOMBRE15" and f2_ws["B47"].value == "NOMBRE30", "F2 perdió el bloque 15–30")
        require(f2_ws["G11"].value == "DOC0001" and f2_ws["G11"].number_format == "@", "F2 no conservó documento como texto")
        require(f2_ws["H11"].value == "2026-09-30" and f2_ws["I11"].value == "LOTE-1" and f2_ws["J11"].value == 2, "F2 no incorporó el hecho confirmado")
        require(f2_ws["Q11"].value == "RECEPTOR DOC-R" and f2_ws["R11"].value == "MADRE" and f2_ws["S11"].value in (None, ""), "F2 alteró receptor, parentesco o firma")
        for row in F2_ROWS[1:]:
            require(all(f2_ws.cell(row, column).value in (None, "") for column in range(8, 20)), "F2 inventó entrega, lote, cantidad, receptor, parentesco o firma")


def main():
    test_order_and_pagination()
    test_real_templates()
    print("PASS test_programas_icbf_phase2")


if __name__ == "__main__":
    main()
