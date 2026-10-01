from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from modules.componente_psicosocial.f23_service import F23Service
from modules.dbapi_compat import sqlite3


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    init_source = (Path(__file__).parents[1] / "init_hosting.py").read_text(encoding="utf-8")
    repo_source = (Path(__file__).parents[1] / "modules" / "componente_psicosocial" / "repository.py").read_text(encoding="utf-8")
    require("F23Service(config_class.DATABASE_PATH, config_class.OUTPUT_FOLDER).init_schema()" in init_source, "Producción no prepara el esquema F23 antes de bloquear DDL")
    require("f.fundacion_id=p.fundacion_id" in repo_source and "p2.fundacion_id=a.fundacion_id" in repo_source, "Los JOIN psicosociales no aíslan ambas tablas por fundación")
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        database = root / "f23.sqlite3"
        with sqlite3.connect(str(database)) as conn:
            conn.executescript("""
            CREATE TABLE master_ninos(
              id INTEGER PRIMARY KEY, fundacion_id INTEGER, activo INTEGER,
              documento TEXT,tipo_documento TEXT,nombres TEXT,apellidos TEXT,
              nombre_completo TEXT,fecha_nacimiento TEXT,sexo TEXT,
              unidad_servicio TEXT,codigo_unidad TEXT,modalidad TEXT,
              vacunas TEXT,control_crecimiento TEXT,peso REAL,talla REAL,datos_json TEXT
            );
            """)
            conn.execute("INSERT INTO master_ninos VALUES(1,1,1,'001234','RC','ANA','PRUEBA','ANA PRUEBA','2022-03-02','F','UDS PRUEBA','00077','INSTITUCIONAL','SI','SI',15.5,96.2,'{}')")
            conn.execute("INSERT INTO master_ninos VALUES(2,1,1,'009999','RC','LUIS','PRUEBA','LUIS PRUEBA',NULL,'M','UDS PRUEBA','00077','INSTITUCIONAL',NULL,NULL,NULL,NULL,'{}')")
            conn.execute("INSERT INTO master_ninos VALUES(3,2,1,'OTRO','RC','OTRA','FUNDACION','OTRA FUNDACION','2021-01-01','F','UDS PRUEBA','00077','INSTITUCIONAL',NULL,NULL,NULL,NULL,'{}')")
            conn.commit()
        service = F23Service(str(database), str(root / "outputs"))
        service.init_schema()
        config = service.configuration(1)
        require(config["template"]["disponible"] and config["template"]["sha256"] == "5340d3e6bfa9f03a09b1400d89c74d5120a81f76a4aa4ea65bae72688dd64319", "No preservó la plantilla oficial")
        require(config["mapeo"]["campos"] == 283, "No cargó los 283 destinos mapeados")
        catalog = {item["id"]: item for item in service.field_catalog()}
        require(catalog["NN-021"]["control"] == "select" and catalog["NN-021"]["options"] == ["SI", "NO"], "Discapacidad no usa respuesta cerrada")
        require(catalog["NN-022"]["depends_on"] == "NN-021" and catalog["NN-022"]["show_when"] == ["SI"], "Categoría de discapacidad no depende de la respuesta principal")
        require(catalog["NN-016"]["control"] == "select" and len(catalog["NN-016"]["options"]) >= 5, "Tipo de documento no ofrece las opciones oficiales iniciales")
        ui_catalogs = service.ui_catalogs()
        require(catalog["NN-002"]["catalog_key"] == "regionales" and "CHOCÓ" in ui_catalogs["regional_centros"], "Regional no ofrece el catálogo territorial")
        require(catalog["NN-003"]["depends_on"] == "NN-002" and len(ui_catalogs["regional_centros"]["CHOCÓ"]) > 0, "Centro zonal no depende de regional")
        require(catalog["NN-014"]["depends_on"] == "NN-013" and len(ui_catalogs["departamento_municipios"]["CHOCÓ"]) > 0, "Municipio no depende del departamento")
        require(len(ui_catalogs["departamento_municipios"]) == 33 and len(ui_catalogs["lenguas"]) > 20, "Los catálogos oficiales quedaron incompletos")
        session = service.create_session(1, "UDS PRUEBA", "2026-09-30", 7)
        require(session["total_participantes"] == 2 and len(session["participantes"]) == 2, "Mezcló otra fundación o perdió participantes")
        ana = next(item for item in session["participantes"] if item["documento"] == "001234")
        require(ana["respuestas"].get("NN-009") == "4 años, 6 meses", "No calculó la edad a la fecha de caracterización")
        require(service.participant_context(1, ana["id"])["unidad"] == "UDS PRUEBA", "No resolvió el ámbito de unidad del participante")
        try:
            service.participant_context(2, ana["id"])
            raise AssertionError("Otra fundación accedió al participante")
        except LookupError:
            pass
        require(ana["pendientes_total"] > 0, "Inventó respuestas para completar la ficha")
        first_pending = next(field_id for field_id in ana["pendientes"] if ana["respuestas"].get(field_id) in (None, ""))
        updated = service.update_participant(1, ana["id"], {"respuestas": {first_pending: "RESPUESTA CONFIRMADA"}, "estados": {first_pending: "CONFIRMADO"}}, 7)
        require(updated["pendientes_total"] == ana["pendientes_total"] - 1, "El guardado parcial no actualizó pendientes")
        second_session = service.create_session(1, "UDS PRUEBA", "2026-10-01", 7)
        second_ana = next(item for item in second_session["participantes"] if item["documento"] == "001234")
        require(second_ana["respuestas"].get(first_pending) == "RESPUESTA CONFIRMADA", "No recuperó la respuesta confirmada de la ficha anterior")
        require(second_ana["estados"].get(first_pending) == "ANTERIOR_POR_CONFIRMAR", "La respuesta anterior se dio por vigente sin confirmación humana")
        require(second_ana["recuperados_total"] >= 1, "No informó los campos recuperados")
        require(next(item for item in second_session["participantes"] if item["documento"] == "009999")["recuperados_total"] == 0, "Copió respuestas a otro participante")
        confirmed_again = service.update_participant(1, second_ana["id"], {"respuestas": {first_pending: "RESPUESTA CONFIRMADA"}, "estados": {first_pending: "CONFIRMADO"}}, 7)
        require(confirmed_again["recuperados_total"] == 0 and confirmed_again["estados"][first_pending] == "CONFIRMADO", "No permitió confirmar la información anterior")
        blocked = False
        try:
            service.update_participant(1, ana["id"], {"integrantes": [{}] * 11}, 7)
        except ValueError:
            blocked = True
        require(blocked, "Truncó o aceptó silenciosamente más de diez integrantes")
        consolidated = service.generate(1, session["id"], 7)
        individual = service.generate(1, session["id"], 7, ana["id"])
        for generated, expected in ((consolidated, 2), (individual, 1)):
            row = service.generation(1, generated["id"])
            path = Path(row["ruta_archivo"])
            require(path.is_file() and generated["participantes"] == expected, "No generó la salida esperada")
            workbook = load_workbook(path, read_only=True, keep_vba=True)
            require(set(F23Service.REGISTER_SHEETS).issubset(workbook.sheetnames), "La salida perdió hojas oficiales")
            require(str(workbook["BD-M1"]["Q2"].value or "") == "001234", "El documento no llegó al registro como texto")
            if expected == 1:
                require(str(workbook["MODULO 1"]["F18"].value or "") == "001234", "La ficha individual no recibió el documento")
            workbook.close()
            try:
                service.generation(2, generated["id"])
                raise AssertionError("Otra fundación accedió a la generación")
            except LookupError:
                pass
            with zipfile.ZipFile(service.template) as original, zipfile.ZipFile(path) as result:
                require(set(original.namelist()) == set(result.namelist()), "La salida perdió componentes del XLSM")
                for name in original.namelist():
                    if not name.startswith("xl/worksheets/"):
                        require(original.read(name) == result.read(name), f"La salida alteró el componente protegido {name}")
        print("PASS test_caracterizacion_f23_v1")


if __name__ == "__main__":
    main()
