from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

RFPP_SHEET = "Formato_asist y entrega RFPP"
F2_SHEET = "Formato_Ajustado_Oficio"
RFPP_ROWS = tuple(range(14, 34))
F2_ROWS = tuple(range(11, 25)) + tuple(range(32, 48))


def safe_name(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("._-")
    return text[:80] or "comunidad"


def spanish_key(value: str) -> tuple:
    text = unicodedata.normalize("NFC", str(value or "").strip().upper())
    alphabet = "ABCDEFGHIJKLMNÑOPQRSTUVWXYZ0123456789"
    ranks = {char: index for index, char in enumerate(alphabet)}
    result = []
    for char in text:
        if char == "Ñ":
            base = char
        else:
            base = "".join(c for c in unicodedata.normalize("NFD", char) if not unicodedata.combining(c))
        for item in base:
            result.append((ranks.get(item, len(ranks)), item))
    return tuple(result)


def participant_sort_key(item: dict, criterion: str) -> tuple:
    data = item.get("canonical") or {}
    names = {
        "pn": data.get("participante.primer_nombre") or "",
        "sn": data.get("participante.segundo_nombre") or "",
        "pa": data.get("participante.primer_apellido") or "",
        "sa": data.get("participante.segundo_apellido") or "",
    }
    order = ("pa", "sa", "pn", "sn") if criterion == "PRIMER_APELLIDO" else ("pn", "sn", "pa", "sa")
    return tuple(spanish_key(names[key]) for key in order) + (int(item.get("id") or 0),)


def sort_participants(items: list[dict], criterion: str | None) -> list[dict]:
    if criterion not in {"PRIMER_APELLIDO", "PRIMER_NOMBRE"}:
        raise ValueError("Confirma primero si el orden principal será por primer apellido o primer nombre.")
    return sorted(items, key=lambda item: participant_sort_key(item, criterion))


def chunks(items: list[dict], size: int) -> list[list[dict]]:
    if not items:
        raise ValueError("La comunidad seleccionada no contiene participantes asignados.")
    return [items[index:index + size] for index in range(0, len(items), size)]


def header_key(item: dict) -> tuple:
    data = item.get("canonical") or {}
    fields = (
        "entidad_contratista.nombre", "contrato.numero", "regional.nombre",
        "municipio.nombre", "centro_zonal.nombre", "unidad.codigo", "unidad.nombre", "unidad.modalidad",
    )
    return tuple(str(data.get(field) or "").strip() for field in fields)


def partition_compatible(items: list[dict]) -> list[list[dict]]:
    groups = {}
    for item in items:
        groups.setdefault(header_key(item), []).append(item)
    return list(groups.values())


def template_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _text(data: dict, field: str) -> str:
    value = data.get(field)
    return "" if value is None else str(value).strip()


def _full_name(data: dict) -> str:
    parts = [_text(data, field) for field in ("participante.primer_nombre", "participante.segundo_nombre", "participante.primer_apellido", "participante.segundo_apellido")]
    return " ".join(part for part in parts if part) or _text(data, "participante.nombre_completo")


def _set_text(cell, value: str) -> None:
    cell.value = value or None
    cell.number_format = "@"


def _clear_cells(ws, rows, columns) -> None:
    for row in rows:
        for column in columns:
            ws.cell(row, column).value = None


def _rfpp_headers(ws, header: dict) -> None:
    labels = {
        "H6": ("Regional:", header.get("regional")),
        "S6": ("Centro Zonal:", header.get("centro_zonal")),
        "A7": ("Municipio:", header.get("municipio")),
        "L7": ("Nombre del operador:", header.get("operador")),
        "L8": ("Nombre del la UDS:", header.get("unidad_nombre")),
    }
    for coordinate, (label, value) in labels.items():
        ws[coordinate] = f"{label} {str(value).strip()}" if value not in (None, "") else label
    optional = {
        "A8": ("Dirección/Lugar del encuentro familiar comunitario:", header.get("lugar")),
        "A9": ("Nombre de los profesionales que lideran el Fortalecimiento Familiar y entregan la RFPP:", header.get("profesionales")),
        "U7": ("Mes del encuentro familiar comunitario y entrega de la RFPP:", header.get("mes")),
        "U9": ("Fecha en la que se realiza el Fortalecimiento Familiar Comunitario:", header.get("fecha_encuentro")),
    }
    for coordinate, (label, value) in optional.items():
        ws[coordinate] = f"{label} {str(value).strip()}" if value not in (None, "") else label


def generate_rfpp(template_path: str | Path, output_path: str | Path, participants: list[dict], header: dict | None = None, facts: dict[int, dict] | None = None) -> dict:
    if len(participants) > len(RFPP_ROWS):
        raise ValueError("Un ejemplar RFPP admite máximo 20 participantes.")
    workbook = load_workbook(template_path, data_only=False, keep_links=True)
    if RFPP_SHEET not in workbook.sheetnames:
        raise ValueError("La plantilla RFPP no contiene la hoja oficial esperada.")
    ws = workbook[RFPP_SHEET]
    _clear_cells(ws, RFPP_ROWS, (2, 6, 7, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22))
    _rfpp_headers(ws, header or {})
    for row, item in zip(RFPP_ROWS, participants):
        data = item.get("canonical") or {}
        _set_text(ws.cell(row, 2), _full_name(data))
        _set_text(ws.cell(row, 6), _text(data, "participante.tipo_documento"))
        _set_text(ws.cell(row, 7), _text(data, "participante.numero_documento"))
        _set_text(ws.cell(row, 19), _text(data, "acudiente.nombre_completo"))
        _set_text(ws.cell(row, 20), _text(data, "acudiente.tipo_documento"))
        _set_text(ws.cell(row, 21), _text(data, "acudiente.numero_documento"))
        fact = (facts or {}).get(int(item.get("id") or 0))
        if fact:
            if fact.get("modalidad") == "COMUNITARIA":
                ws.cell(row, 12).value = "SI"
            elif fact.get("modalidad") == "HOGAR":
                _set_text(ws.cell(row, 14), str(fact.get("fecha_entrega") or ""))
                ws.cell(row, 17).value = "SI"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {"path": str(output_path), "format": "RFPP", "participants": len(participants), "template_sha256": template_sha256(template_path)}


def _f2_headers(ws, header: dict) -> None:
    destinations = {
        "G2": "punto_codigo", "O2": "mes", "S2": "anio",
        "C3": "regional", "G3": "punto_nombre", "C4": "centro_zonal",
        "G4": "responsable", "O4": "suplente", "C5": "municipio",
        "G5": "direccion", "O5": "barrio", "S5": "telefono",
        "C6": "modalidad", "G6": "punto_origen_codigo", "O6": "punto_origen_nombre",
    }
    for coordinate, field in destinations.items():
        _set_text(ws[coordinate], str(header.get(field) or "").strip())


def generate_f2(template_path: str | Path, output_path: str | Path, participants: list[dict], header: dict | None = None, facts: dict[int, dict] | None = None) -> dict:
    if len(participants) > len(F2_ROWS):
        raise ValueError("Un ejemplar F2 admite máximo 30 participantes.")
    workbook = load_workbook(template_path, data_only=False, keep_links=True)
    if F2_SHEET not in workbook.sheetnames:
        raise ValueError("La plantilla F2 no contiene la hoja oficial esperada.")
    ws = workbook[F2_SHEET]
    _clear_cells(ws, F2_ROWS, range(2, 20))
    _f2_headers(ws, header or {})
    for row, item in zip(F2_ROWS, participants):
        data = item.get("canonical") or {}
        values = (
            _text(data, "participante.primer_nombre"), _text(data, "participante.segundo_nombre"),
            _text(data, "participante.primer_apellido"), _text(data, "participante.segundo_apellido"),
            _text(data, "participante.tipo_documento"), _text(data, "participante.numero_documento"),
        )
        for column, value in zip(range(2, 8), values):
            _set_text(ws.cell(row, column), value)
        fact = (facts or {}).get(int(item.get("id") or 0))
        if fact:
            _set_text(ws.cell(row, 8), str(fact.get("fecha_entrega") or ""))
            for product in fact.get("items") or []:
                normalized = normalize_product(product.get("producto"))
                columns = {"BIENESTARINA_MAS": (9, 10), "BIENESTARINA_LIQUIDA": (11, 12), "ALIMENTO_MG_ML": (13, 14), "OTRO": (15, 16)}[normalized]
                _set_text(ws.cell(row, columns[0]), str(product.get("lote") or ""))
                ws.cell(row, columns[1]).value = product.get("unidades")
            receiver = " ".join(value for value in (str(fact.get("receptor_nombre") or "").strip(), str(fact.get("receptor_documento") or "").strip()) if value)
            _set_text(ws.cell(row, 17), receiver)
            _set_text(ws.cell(row, 18), str(fact.get("receptor_parentesco") or ""))
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {"path": str(output_path), "format": "F2", "participants": len(participants), "template_sha256": template_sha256(template_path)}


def normalize_product(value: str | None) -> str:
    text = safe_name(str(value or "")).replace("_", " ").upper()
    if "BIENESTARINA" in text and ("LIQUID" in text):
        return "BIENESTARINA_LIQUIDA"
    if "BIENESTARINA" in text and ("MAS" in text or "MÁS" in str(value or "").upper()):
        return "BIENESTARINA_MAS"
    if ("MG" in text and "ML" in text) or "GESTACION" in text or "LACTANCIA" in text:
        return "ALIMENTO_MG_ML"
    return "OTRO"


def generate_pages(format_code: str, template_path: str | Path, output_dir: str | Path, community_name: str, participants: list[dict], criterion: str | None, header: dict | None = None, facts: dict[int, dict] | None = None) -> list[dict]:
    ordered = sort_participants(participants, criterion)
    output_dir = Path(output_dir)
    results = []
    capacity = 20 if format_code == "RFPP" else 30
    generator = generate_rfpp if format_code == "RFPP" else generate_f2 if format_code == "F2" else None
    if generator is None:
        raise ValueError("Formato no soportado para este perfil.")
    compatible_groups = partition_compatible(ordered)
    for group_number, group in enumerate(compatible_groups, 1):
        data = group[0].get("canonical") or {}
        derived = {
            "regional": data.get("regional.nombre"),
            "centro_zonal": data.get("centro_zonal.nombre"),
            "municipio": data.get("municipio.nombre"),
            "operador": data.get("entidad_contratista.nombre"),
            "unidad_nombre": data.get("unidad.nombre"),
            "punto_codigo": data.get("unidad.codigo"),
            "punto_nombre": data.get("unidad.nombre"),
            "modalidad": data.get("unidad.modalidad") or data.get("programa.modalidad"),
        }
        effective_header = {**derived, **{key: value for key, value in (header or {}).items() if value not in (None, "")}}
        for page_number, page in enumerate(chunks(group, capacity), 1):
            name = f"{format_code}_{safe_name(community_name)}_g{group_number:02d}_p{page_number:02d}.xlsx"
            result = generator(template_path, output_dir / name, page, effective_header, facts or {})
            result.update({"group": group_number, "page": page_number, "filename": name, "participant_ids": [int(item.get("id") or 0) for item in page]})
            results.append(result)
    return results
