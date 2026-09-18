"""ALPHA51: Gestión de entregables del componente Salud y Nutrición.

Este módulo es deliberadamente independiente del motor Pack35 de carga,
CoreCursor y formatos oficiales. Usa el repositorio ya registrado para crear
un catálogo mensual, evidencias, documentos de soporte y paquete ZIP sin tocar
la lógica estable de procesamiento.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from werkzeug.utils import secure_filename

from modules.seguridad.tenant_context import current_tenant_context, current_tenant_id

from .services import now_iso

ALLOWED_EVIDENCE_EXT = {'.png', '.jpg', '.jpeg', '.webp', '.pdf', '.doc', '.docx', '.xlsx', '.xls'}

ENTREGABLES_SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS sn_entregables_catalogo (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    fundacion_id INTEGER,
    codigo TEXT NOT NULL UNIQUE,
    nombre TEXT NOT NULL,
    descripcion TEXT,
    actividad TEXT,
    evidencias_requeridas TEXT,
    forma_entrega TEXT DEFAULT 'Físico',
    responsable TEXT DEFAULT 'Enfermera/Nutricionista',
    periodicidad TEXT DEFAULT 'Mensual',
    aplica_por_uds INTEGER DEFAULT 1,
    requiere_acta INTEGER DEFAULT 0,
    requiere_listado INTEGER DEFAULT 0,
    requiere_fotos INTEGER DEFAULT 0,
    minimo_fotos INTEGER DEFAULT 0,
    requiere_oficio INTEGER DEFAULT 0,
    requiere_formato_excel INTEGER DEFAULT 0,
    requiere_word INTEGER DEFAULT 0,
    requiere_pdf INTEGER DEFAULT 0,
    requiere_firma INTEGER DEFAULT 0,
    plantilla_asociada TEXT,
    estado TEXT DEFAULT 'ACTIVO',
    observaciones TEXT,
    fecha_creacion TEXT,
    fecha_actualizacion TEXT
);

CREATE TABLE IF NOT EXISTS sn_entregables_mes (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    catalogo_id INTEGER NOT NULL,
    codigo TEXT NOT NULL,
    mes INTEGER NOT NULL,
    anio INTEGER NOT NULL,
    uds TEXT DEFAULT 'TODAS',
    coordinador TEXT,
    fundacion_id INTEGER DEFAULT 1,
    corporacion_id INTEGER DEFAULT 1,
    responsable TEXT,
    estado TEXT DEFAULT 'pendiente',
    observaciones TEXT,
    porcentaje INTEGER DEFAULT 0,
    fecha_creacion TEXT,
    fecha_actualizacion TEXT,
    usuario_id INTEGER,
    UNIQUE(catalogo_id, mes, anio, uds, fundacion_id),
    FOREIGN KEY (catalogo_id) REFERENCES sn_entregables_catalogo(id)
);

CREATE TABLE IF NOT EXISTS sn_entregables_evidencias (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    entregable_id INTEGER NOT NULL,
    fundacion_id INTEGER DEFAULT 1,
    nombre_original TEXT,
    nombre_guardado TEXT,
    ruta_archivo TEXT NOT NULL,
    tipo TEXT DEFAULT 'foto',
    actividad TEXT,
    fecha_actividad TEXT,
    uds TEXT,
    responsable TEXT,
    observaciones TEXT,
    fecha_carga TEXT,
    usuario_id INTEGER,
    FOREIGN KEY (entregable_id) REFERENCES sn_entregables_mes(id)
);

CREATE TABLE IF NOT EXISTS sn_entregables_archivos (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    entregable_id INTEGER,
    fundacion_id INTEGER DEFAULT 1,
    tipo TEXT NOT NULL,
    nombre_archivo TEXT NOT NULL,
    ruta_archivo TEXT NOT NULL,
    estado TEXT DEFAULT 'generado',
    metadata_json TEXT,
    fecha_generacion TEXT,
    usuario_id INTEGER,
    FOREIGN KEY (entregable_id) REFERENCES sn_entregables_mes(id)
);

CREATE TABLE IF NOT EXISTS sn_entregables_validaciones (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    entregable_id INTEGER NOT NULL,
    fundacion_id INTEGER DEFAULT 1,
    valido INTEGER DEFAULT 0,
    pendientes_json TEXT,
    resultado_json TEXT,
    fecha_validacion TEXT,
    usuario_id INTEGER,
    FOREIGN KEY (entregable_id) REFERENCES sn_entregables_mes(id)
);

CREATE TABLE IF NOT EXISTS sn_entregables_observaciones (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    entregable_id INTEGER NOT NULL,
    fundacion_id INTEGER DEFAULT 1,
    observacion TEXT NOT NULL,
    estado TEXT DEFAULT 'abierta',
    fecha_creacion TEXT,
    usuario_id INTEGER,
    FOREIGN KEY (entregable_id) REFERENCES sn_entregables_mes(id)
);

CREATE TABLE IF NOT EXISTS sn_entregables_actividades (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    entregable_id INTEGER NOT NULL,
    fundacion_id INTEGER DEFAULT 1,
    fecha_actividad TEXT NOT NULL,
    hora_inicio TEXT,
    hora_final TEXT,
    lugar TEXT,
    dirigido_a TEXT,
    objetivo TEXT,
    agenda_json TEXT,
    desarrollo TEXT,
    resultados TEXT,
    compromisos TEXT,
    dificultades TEXT,
    acciones_mejora TEXT,
    participantes_total INTEGER DEFAULT 0,
    responsable TEXT,
    confirmado INTEGER DEFAULT 0,
    fecha_confirmacion TEXT,
    fecha_creacion TEXT,
    fecha_actualizacion TEXT,
    usuario_id INTEGER,
    FOREIGN KEY (entregable_id) REFERENCES sn_entregables_mes(id)
);

CREATE TABLE IF NOT EXISTS sn_entregables_plantillas_oficiales (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    fundacion_id INTEGER DEFAULT 1,
    codigo_entregable TEXT NOT NULL,
    tipo TEXT NOT NULL,
    nombre_original TEXT NOT NULL,
    ruta_archivo TEXT NOT NULL,
    extension TEXT NOT NULL,
    sha256 TEXT NOT NULL,
    version INTEGER DEFAULT 1,
    estado TEXT DEFAULT 'ACTIVA',
    fecha_registro TEXT,
    usuario_id INTEGER
);

CREATE INDEX IF NOT EXISTS idx_sn_entregables_mes_periodo ON sn_entregables_mes(anio, mes, uds);
CREATE INDEX IF NOT EXISTS idx_sn_entregables_mes_estado ON sn_entregables_mes(estado);
CREATE INDEX IF NOT EXISTS idx_sn_entregables_evidencias_entregable ON sn_entregables_evidencias(entregable_id);
CREATE INDEX IF NOT EXISTS idx_sn_entregables_archivos_entregable ON sn_entregables_archivos(entregable_id);
CREATE INDEX IF NOT EXISTS idx_sn_entregables_actividades_entregable ON sn_entregables_actividades(entregable_id);
CREATE INDEX IF NOT EXISTS idx_sn_entregables_plantillas_codigo ON sn_entregables_plantillas_oficiales(codigo_entregable, tipo, fundacion_id, estado);
"""

CATALOGO_BASE = [
    dict(codigo='E01_REGISTRO_NOVEDADES', nombre='Registro de novedades, articulaciones en salud y garantía de derechos',
         actividad='Acta de revisión de carpeta; identificación de garantías de derechos y canalizaciones.',
         evidencias='Acta de revisión, formato de novedades, oficio/correo de articulación y evidencias fotográficas.', acta=1, listado=0, fotos=1, oficio=1, excel=1, word=1),
    dict(codigo='E02_LAVADO_MANOS', nombre='Lavado de manos',
         actividad='Promover con usuarios y familias la práctica correcta del lavado de manos.',
         evidencias='Acta, listado de asistencia y evidencias fotográficas.', acta=1, listado=1, fotos=1, oficio=0, excel=0, word=1),
    dict(codigo='E03_LACTANCIA', nombre='Lactancia materna / extracción de lactancia',
         actividad='Promover práctica de lactancia con gestantes, madres lactantes y familias.',
         evidencias='Acta, listado de asistencia y evidencias fotográficas.', acta=1, listado=1, fotos=1, oficio=0, excel=0, word=1),
    dict(codigo='E04_FICHA_LACTANCIA', nombre='Ficha de observación de lactancia materna',
         actividad='Aplicación y diligenciamiento de formato de lactancia materna.',
         evidencias='Formato ficha de observación y soportes fotográficos.', acta=0, listado=0, fotos=1, oficio=0, excel=1, word=0),
    dict(codigo='E05_ANTROPOMETRIA', nombre='Formato de captura de medidas antropométricas',
         actividad='Toma de peso, talla y perímetro braquial; usuarios nuevos, gestantes y seguimiento a riesgo/DNT.',
         evidencias='Formato de captura por grupo, reporte de Cuéntame y fotografías de valoración.', acta=0, listado=0, fotos=1, oficio=0, excel=1, word=0),
    dict(codigo='E06_SIGNOS_FISICOS', nombre='Formato de identificación de signos físicos',
         actividad='Aplicación y diligenciamiento de formato de signos físicos.',
         evidencias='Formato Excel magnético y muestra física en informe.', acta=0, listado=0, fotos=0, oficio=0, excel=1, word=0),
    dict(codigo='E07_MONITOREO_DNT', nombre='Formato de monitoreo de signos de alarma DNT',
         actividad='Monitoreo semanal para DNT, riesgo de DNT y signos físicos.',
         evidencias='Formato Excel magnético y anexo físico en informe.', acta=0, listado=0, fotos=0, oficio=0, excel=1, word=0),
    dict(codigo='E08_CONCERTACION_MINUTA', nombre='Concertación de minuta RPP',
         actividad='Concertar la minuta con las familias beneficiarias de la modalidad.',
         evidencias='Acta, listado de asistencia y evidencias fotográficas.', acta=1, listado=1, fotos=1, oficio=0, excel=0, word=1),
    dict(codigo='E09_CONTROL_CALIDAD', nombre='Control de calidad de raciones y alimentos',
         actividad='Verificación de raciones alimentarias, almacenamiento y entrega de RPP/Bienestarina.',
         evidencias='Acta de verificación de alimentos y evidencias de verificación/almacenamiento.', acta=1, listado=0, fotos=1, oficio=0, excel=1, word=1),
    dict(codigo='E10_LIMPIEZA_DESINFECCION', nombre='Acta de limpieza y desinfección de UCAS',
         actividad='Verificación de limpieza y desinfección en UCAS.',
         evidencias='Acta, evidencias fotográficas y muestras de saneamiento básico.', acta=1, listado=0, fotos=1, oficio=0, excel=1, word=1),
    dict(codigo='E11_ENCUENTROS_HOGAR', nombre='Encuentros en el hogar',
         actividad='Encuentros en hogares priorizados: puerperio, DNT, riesgo, enfermedad o garantía de derechos.',
         evidencias='Formato de encuentro en el hogar, mínimo 2 muestras por informe y evidencias fotográficas.', acta=0, listado=1, fotos=1, oficio=0, excel=1, word=1),
    dict(codigo='E12_ARTICULACION_INTERINSTITUCIONAL', nombre='Articulación interinstitucional',
         actividad='Oficio de acuerdo a la situación encontrada.',
         evidencias='Oficio físico/digital.', acta=0, listado=0, fotos=0, oficio=1, excel=0, word=1),
    dict(codigo='E13_ENTREGA_RPP', nombre='Entrega de RPP',
         actividad='Evidencias fotográficas de entrega de RPP.',
         evidencias='Evidencias fotográficas de entrega de RPP.', acta=0, listado=0, fotos=1, oficio=0, excel=0, word=0),
    dict(codigo='E14_OLLAS_REFRIGERIOS', nombre='Preparaciones de ollas comunitarias o entrega de refrigerios',
         actividad='Evidencias de preparación de olla comunitaria o entrega/consumo de refrigerios.',
         evidencias='Evidencias fotográficas con participación de familias.', acta=0, listado=1, fotos=1, oficio=0, excel=0, word=1),
    dict(codigo='E15_CUALIFICACION_TH', nombre='Acta de cualificación al talento humano',
         actividad='Cualificación en lactancia, almacenamiento de alimentos y control de plagas.',
         evidencias='Acta, listado de asistencia y evidencias fotográficas.', acta=1, listado=1, fotos=1, oficio=0, excel=0, word=1),
    dict(codigo='E16_NOVEDADES_DNT_ETA', nombre='Novedades, canalización y certificado de no ETA',
         actividad='Registrar novedades por malnutrición, perímetro braquial menor de 11.5 cm, signos clínicos o enfermedad; activar ruta y elaborar certificado cuando aplique.',
         evidencias='Formato oficial de novedades, oficio/correo de canalización y certificado de no ETA.', acta=0, listado=0, fotos=0, oficio=1, excel=0, word=1),
]


def _slug(text: str) -> str:
    text = str(text or '').upper()
    text = re.sub(r'[^A-Z0-9]+', '_', text)
    return text.strip('_') or 'SIN_DATO'


def _month_name(mes: int) -> str:
    meses = ['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']
    return meses[mes - 1] if 1 <= mes <= 12 else str(mes)


def _split_uds(value: Any) -> list[str]:
    if isinstance(value, list):
        raw = value
    else:
        raw = re.split(r'[,|;\n]+', str(value or ''))
    uds = []
    seen = set()
    for item in raw:
        cleaned = str(item or '').strip()
        if cleaned and cleaned.upper() not in seen:
            seen.add(cleaned.upper())
            uds.append(cleaned)
    return uds or ['TODAS']


class EntregablesSaludNutricionService:
    def __init__(self, repo, output_folder: str, upload_folder: str):
        self.repo = repo
        self._output_folder = output_folder
        self._upload_folder = upload_folder

    @property
    def output_folder(self) -> Path:
        path = Path(os.fspath(self._output_folder)) / 'salud_nutricion' / 'entregables'
        path.mkdir(parents=True, exist_ok=True)
        return path

    @property
    def upload_folder(self) -> Path:
        path = Path(os.fspath(self._upload_folder)) / 'salud_nutricion' / 'entregables'
        path.mkdir(parents=True, exist_ok=True)
        return path

    @property
    def template_folder(self) -> Path:
        path = self.upload_folder / 'plantillas_oficiales'
        path.mkdir(parents=True, exist_ok=True)
        return path

    def init_schema(self) -> None:
        self.repo.execute_script(ENTREGABLES_SCHEMA_SQL)
        for table in (
            'sn_entregables_catalogo', 'sn_entregables_mes', 'sn_entregables_evidencias',
            'sn_entregables_archivos', 'sn_entregables_validaciones',
            'sn_entregables_observaciones', 'sn_entregables_actividades',
            'sn_entregables_plantillas_oficiales',
        ):
            if self.repo.table_exists(table):
                default = 'INTEGER' if table == 'sn_entregables_catalogo' else 'INTEGER DEFAULT 1'
                self.repo.ensure_column(table, 'fundacion_id', default)
        self.seed_catalogo()

    def seed_catalogo(self) -> None:
        now = now_iso()
        # Migración reversible de los dos nombres provisionales usados antes de
        # recibir la matriz institucional de agosto de 2026. Se ejecuta antes
        # del alta para conservar IDs y evitar duplicados en bases existentes.
        aliases = {
            'E02_AAVN': CATALOGO_BASE[1],
            'E08_SOCIALIZACION_RPP': CATALOGO_BASE[7],
        }
        for old_code, item in aliases.items():
            old = self.repo.fetch_one('SELECT id FROM sn_entregables_catalogo WHERE codigo = ?', (old_code,))
            canonical = self.repo.fetch_one('SELECT id FROM sn_entregables_catalogo WHERE codigo = ?', (item['codigo'],))
            if old and not canonical:
                self.repo.execute(
                    'UPDATE sn_entregables_catalogo SET codigo = ? WHERE id = ?',
                    (item['codigo'], old['id']),
                )
                self.repo.execute(
                    'UPDATE sn_entregables_mes SET codigo = ? WHERE catalogo_id = ?',
                    (item['codigo'], old['id']),
                )

        for idx, item in enumerate(CATALOGO_BASE, start=1):
            found = self.repo.fetch_one('SELECT id FROM sn_entregables_catalogo WHERE codigo = ?', (item['codigo'],))
            if found:
                self._sync_catalog_item(int(found['id']), item, now)
                continue
            self.repo.execute(
                """
                INSERT INTO sn_entregables_catalogo
                (codigo, nombre, descripcion, actividad, evidencias_requeridas, forma_entrega, responsable,
                 periodicidad, aplica_por_uds, requiere_acta, requiere_listado, requiere_fotos, minimo_fotos,
                 requiere_oficio, requiere_formato_excel, requiere_word, requiere_pdf, requiere_firma,
                 plantilla_asociada, estado, observaciones, fecha_creacion, fecha_actualizacion)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    item['codigo'], item['nombre'], item.get('descripcion') or item['nombre'], item['actividad'],
                    item['evidencias'], 'Físico', 'Enfermera/Nutricionista', 'Mensual', 1,
                    item.get('acta', 0), item.get('listado', 0), item.get('fotos', 0), 4 if item.get('fotos', 0) else 0,
                    item.get('oficio', 0), item.get('excel', 0), item.get('word', 0), 0,
                    1 if item.get('acta') or item.get('listado') else 0,
                    item['codigo'], 'ACTIVO', 'Catálogo base ALPHA51', now, now,
                ),
            )

    def _sync_catalog_item(self, catalog_id: int, item: dict[str, Any], now: str) -> None:
        self.repo.execute(
            """
            UPDATE sn_entregables_catalogo
               SET nombre = ?, descripcion = ?, actividad = ?, evidencias_requeridas = ?,
                   requiere_acta = ?, requiere_listado = ?, requiere_fotos = ?, minimo_fotos = ?,
                   requiere_oficio = ?, requiere_formato_excel = ?, requiere_word = ?,
                   requiere_firma = ?, plantilla_asociada = ?, fecha_actualizacion = ?
             WHERE id = ?
            """,
            (
                item['nombre'], item.get('descripcion') or item['nombre'], item['actividad'], item['evidencias'],
                item.get('acta', 0), item.get('listado', 0), item.get('fotos', 0),
                4 if item.get('fotos', 0) else 0, item.get('oficio', 0), item.get('excel', 0),
                item.get('word', 0), 1 if item.get('acta') or item.get('listado') else 0,
                item['codigo'], now, catalog_id,
            ),
        )

    def catalogo(self) -> list[dict[str, Any]]:
        return self.repo.fetch_all('SELECT * FROM sn_entregables_catalogo ORDER BY id ASC')

    def registrar_plantilla(self, codigo: str, tipo: str, archivo, usuario_id: int | None = None) -> dict[str, Any]:
        codigo = str(codigo or '').strip().upper()
        tipo = str(tipo or '').strip().lower()
        if tipo not in {'acta', 'listado', 'formato', 'oficio', 'informe'}:
            raise ValueError('Tipo de plantilla no permitido.')
        if not self.repo.fetch_one('SELECT id FROM sn_entregables_catalogo WHERE codigo = ?', (codigo,)):
            raise ValueError('El entregable indicado no existe en el catálogo.')
        original = secure_filename(archivo.filename or '')
        extension = Path(original).suffix.lower()
        permitidas = {'.docx'} if tipo in {'acta', 'oficio', 'informe'} else {'.xlsx'}
        if extension not in permitidas:
            esperado = 'DOCX' if '.docx' in permitidas else 'XLSX'
            raise ValueError(f'La plantilla {tipo} debe ser {esperado}.')
        contenido = archivo.read()
        if not contenido:
            raise ValueError('La plantilla está vacía.')
        fundacion_id = int(current_tenant_id(1) or 1)
        digest = hashlib.sha256(contenido).hexdigest()
        version_row = self.repo.fetch_one(
            'SELECT MAX(version) AS version FROM sn_entregables_plantillas_oficiales WHERE codigo_entregable = ? AND tipo = ? AND fundacion_id = ?',
            (codigo, tipo, fundacion_id),
        ) or {}
        version = int(version_row.get('version') or 0) + 1
        destino = self.template_folder / f'{codigo}_{tipo}_v{version}{extension}'
        destino.write_bytes(contenido)
        self.repo.execute(
            "UPDATE sn_entregables_plantillas_oficiales SET estado = 'INACTIVA' WHERE codigo_entregable = ? AND tipo = ? AND fundacion_id = ?",
            (codigo, tipo, fundacion_id),
        )
        self.repo.execute(
            '''INSERT INTO sn_entregables_plantillas_oficiales
               (fundacion_id, codigo_entregable, tipo, nombre_original, ruta_archivo, extension, sha256, version, estado, fecha_registro, usuario_id)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'ACTIVA', ?, ?)''',
            (fundacion_id, codigo, tipo, original, str(destino), extension, digest, version, now_iso(), usuario_id),
        )
        return self.repo.fetch_one('SELECT * FROM sn_entregables_plantillas_oficiales WHERE ruta_archivo = ?', (str(destino),)) or {}

    def listar_plantillas(self, codigo: str | None = None) -> list[dict[str, Any]]:
        fundacion_id = int(current_tenant_id(1) or 1)
        sql = "SELECT id, codigo_entregable, tipo, nombre_original, extension, sha256, version, estado, fecha_registro FROM sn_entregables_plantillas_oficiales WHERE fundacion_id = ? AND estado = 'ACTIVA'"
        params: list[Any] = [fundacion_id]
        if codigo:
            sql += ' AND codigo_entregable = ?'
            params.append(str(codigo).upper())
        return self.repo.fetch_all(sql + ' ORDER BY codigo_entregable, tipo', tuple(params))

    def _plantilla_activa(self, codigo: str, tipo: str) -> dict[str, Any]:
        row = self.repo.fetch_one(
            "SELECT * FROM sn_entregables_plantillas_oficiales WHERE codigo_entregable = ? AND tipo = ? AND fundacion_id = ? AND estado = 'ACTIVA' ORDER BY version DESC LIMIT 1",
            (codigo, tipo, int(current_tenant_id(1) or 1)),
        )
        if not row or not Path(str(row.get('ruta_archivo') or '')).is_file():
            raise ValueError(f'Falta cargar la plantilla oficial de {tipo} para {codigo}. No se generará un formato genérico.')
        return row

    def crear_mes(self, payload: dict[str, Any]) -> dict[str, Any]:
        mes = int(payload.get('mes') or datetime.now().month)
        anio = int(payload.get('anio') or datetime.now().year)
        uds_list = _split_uds(payload.get('uds') or payload.get('unidades') or 'TODAS')
        coordinador = str(payload.get('coordinador') or '').strip()
        responsable = str(payload.get('responsable') or payload.get('usuario') or 'Enfermera/Nutricionista').strip()
        context = current_tenant_context()
        requested_fundacion = payload.get('fundacion_id')
        if context.role == 'SUPERADMIN' and context.allow_global and requested_fundacion:
            fundacion_id = int(requested_fundacion)
        else:
            fundacion_id = int(context.tenant_id or 1)
        corporacion_id = int(payload.get('corporacion_id') or 1)
        now = now_iso()
        creados = 0
        existentes = 0
        for cat in self.catalogo():
            for uds in uds_list:
                found = self.repo.fetch_one(
                    'SELECT id FROM sn_entregables_mes WHERE catalogo_id = ? AND mes = ? AND anio = ? AND uds = ? AND fundacion_id = ?',
                    (cat['id'], mes, anio, uds, fundacion_id),
                )
                if found:
                    existentes += 1
                    continue
                self.repo.execute(
                    """
                    INSERT INTO sn_entregables_mes
                    (catalogo_id, codigo, mes, anio, uds, coordinador, fundacion_id, corporacion_id,
                     responsable, estado, observaciones, porcentaje, fecha_creacion, fecha_actualizacion, usuario_id)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (cat['id'], cat['codigo'], mes, anio, uds, coordinador, fundacion_id, corporacion_id,
                     responsable, 'pendiente', '', 0, now, now, payload.get('usuario_id')),
                )
                creados += 1
        return {'mes': mes, 'anio': anio, 'uds': uds_list, 'creados': creados, 'existentes': existentes}

    def listar(self, filtros: dict[str, Any]) -> dict[str, Any]:
        fundacion_id = int(current_tenant_id(1) or 1)
        where = [f'COALESCE(m.fundacion_id, 1) = {fundacion_id}']
        params: list[Any] = []
        if filtros.get('mes'):
            where.append('m.mes = ?')
            params.append(int(filtros['mes']))
        if filtros.get('anio'):
            where.append('m.anio = ?')
            params.append(int(filtros['anio']))
        if filtros.get('uds'):
            where.append('m.uds = ?')
            params.append(filtros['uds'])
        if filtros.get('estado'):
            where.append('m.estado = ?')
            params.append(filtros['estado'])
        rows = self.repo.fetch_all(
            f"""
            SELECT m.*, c.nombre, c.actividad, c.evidencias_requeridas, c.requiere_acta, c.requiere_listado,
                   c.requiere_fotos, c.minimo_fotos, c.requiere_oficio, c.requiere_formato_excel,
                   c.requiere_word, c.requiere_pdf, c.plantilla_asociada,
                   (SELECT GROUP_CONCAT(p.tipo, ',') FROM sn_entregables_plantillas_oficiales p
                     WHERE p.codigo_entregable = m.codigo
                       AND p.fundacion_id = {fundacion_id} AND p.estado = 'ACTIVA') AS plantillas_cargadas,
                   (SELECT COUNT(*) FROM sn_entregables_evidencias e
                     WHERE e.entregable_id = m.id
                       AND COALESCE(e.fundacion_id, 1) = {fundacion_id}) AS fotos_cargadas,
                   (SELECT COUNT(*) FROM sn_entregables_archivos a
                     WHERE a.entregable_id = m.id
                       AND COALESCE(a.fundacion_id, 1) = {fundacion_id}) AS archivos_generados
                   ,(SELECT COUNT(*) FROM sn_entregables_actividades x
                     WHERE x.entregable_id = m.id
                       AND COALESCE(x.fundacion_id, 1) = {fundacion_id}
                       AND x.confirmado = 1) AS actividades_confirmadas
            FROM sn_entregables_mes m
            JOIN sn_entregables_catalogo c
              ON c.id = m.catalogo_id
             AND (c.fundacion_id IS NULL OR c.fundacion_id = {fundacion_id})
            WHERE {' AND '.join(where)}
            ORDER BY m.anio DESC, m.mes DESC, m.uds ASC, c.id ASC
            LIMIT 3000
            """,
            params,
        )
        resumen = self._resumen(rows)
        return {'entregables': rows, 'resumen': resumen}

    def detalle(self, entregable_id: int) -> dict[str, Any] | None:
        fundacion_id = int(current_tenant_id(1) or 1)
        row = self.repo.fetch_one(
            f"""
            SELECT m.*, c.nombre, c.actividad, c.evidencias_requeridas, c.requiere_acta, c.requiere_listado,
                   c.requiere_fotos, c.minimo_fotos, c.requiere_oficio, c.requiere_formato_excel,
                   c.requiere_word, c.requiere_pdf, c.plantilla_asociada
            FROM sn_entregables_mes m
            JOIN sn_entregables_catalogo c
              ON c.id = m.catalogo_id
             AND (c.fundacion_id IS NULL OR c.fundacion_id = {fundacion_id})
            WHERE m.id = ?
              AND COALESCE(m.fundacion_id, 1) = {fundacion_id}
            """,
            (entregable_id,),
        )
        if not row:
            return None
        row['evidencias'] = self.repo.fetch_all('SELECT * FROM sn_entregables_evidencias WHERE entregable_id = ? ORDER BY id', (entregable_id,))
        row['archivos'] = self.repo.fetch_all('SELECT * FROM sn_entregables_archivos WHERE entregable_id = ? ORDER BY id DESC', (entregable_id,))
        row['validaciones'] = self.repo.fetch_all('SELECT * FROM sn_entregables_validaciones WHERE entregable_id = ? ORDER BY id DESC LIMIT 10', (entregable_id,))
        row['actividades'] = self.listar_actividades(entregable_id)
        return row

    def listar_actividades(self, entregable_id: int) -> list[dict[str, Any]]:
        fundacion_id = int(current_tenant_id(1) or 1)
        return self.repo.fetch_all(
            """
            SELECT * FROM sn_entregables_actividades
             WHERE entregable_id = ? AND COALESCE(fundacion_id, 1) = ?
             ORDER BY fecha_actividad ASC, id ASC
            """,
            (entregable_id, fundacion_id),
        )

    def guardar_actividad(self, entregable_id: int, payload: dict[str, Any], usuario_id: int | None = None) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        fecha = str(payload.get('fecha_actividad') or '').strip()
        objetivo = str(payload.get('objetivo') or '').strip()
        desarrollo = str(payload.get('desarrollo') or '').strip()
        responsable = str(payload.get('responsable') or ent.get('responsable') or '').strip()
        confirmado = 1 if str(payload.get('confirmado') or '').lower() in {'1', 'true', 'si', 'sí', 'on'} else 0
        if not fecha:
            raise ValueError('La fecha de la actividad es obligatoria.')
        if confirmado and (not objetivo or not desarrollo or not responsable):
            raise ValueError('Para confirmar la actividad se requieren objetivo, desarrollo y responsable.')
        agenda = payload.get('agenda') or []
        if isinstance(agenda, str):
            agenda = [line.strip() for line in agenda.splitlines() if line.strip()]
        now = now_iso()
        activity_id = self.repo.execute(
            """
            INSERT INTO sn_entregables_actividades
            (entregable_id, fundacion_id, fecha_actividad, hora_inicio, hora_final, lugar, dirigido_a,
             objetivo, agenda_json, desarrollo, resultados, compromisos, dificultades, acciones_mejora,
             participantes_total, responsable, confirmado, fecha_confirmacion, fecha_creacion,
             fecha_actualizacion, usuario_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                entregable_id, int(current_tenant_id(1) or 1), fecha,
                str(payload.get('hora_inicio') or '').strip(), str(payload.get('hora_final') or '').strip(),
                str(payload.get('lugar') or ent.get('uds') or '').strip(), str(payload.get('dirigido_a') or '').strip(),
                objetivo, json.dumps(agenda, ensure_ascii=False), desarrollo,
                str(payload.get('resultados') or '').strip(), str(payload.get('compromisos') or '').strip(),
                str(payload.get('dificultades') or '').strip(), str(payload.get('acciones_mejora') or '').strip(),
                max(0, int(payload.get('participantes_total') or 0)), responsable, confirmado,
                now if confirmado else None, now, now, usuario_id,
            ),
        )
        if not activity_id:
            latest = self.repo.fetch_one(
                'SELECT id FROM sn_entregables_actividades WHERE entregable_id = ? ORDER BY id DESC LIMIT 1',
                (entregable_id,),
            ) or {}
            activity_id = int(latest.get('id') or 0)
        self._touch(entregable_id, 'en proceso')
        return {'id': activity_id, 'confirmado': bool(confirmado), 'fecha_actividad': fecha}

    def _resumen(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        total = len(rows)
        completos = sum(1 for r in rows if str(r.get('estado') or '').lower() == 'completo')
        pendientes = sum(1 for r in rows if str(r.get('estado') or '').lower() == 'pendiente')
        observados = sum(1 for r in rows if str(r.get('estado') or '').lower() == 'observado')
        fotos_faltantes = 0
        for r in rows:
            minimo = int(r.get('minimo_fotos') or 0) if int(r.get('requiere_fotos') or 0) else 0
            fotos = int(r.get('fotos_cargadas') or 0)
            fotos_faltantes += max(0, minimo - fotos)
        return {
            'total': total,
            'completos': completos,
            'pendientes': pendientes,
            'observados': observados,
            'fotos_faltantes': fotos_faltantes,
            'porcentaje': round((completos / total) * 100, 1) if total else 0,
            'actividades_confirmadas': sum(int(r.get('actividades_confirmadas') or 0) for r in rows),
        }

    def obtener_usuarios_base(self, uds: str | None = None, limit: int = 2000) -> list[dict[str, Any]]:
        # Fuente poblacional única: Base Maestra publicada.
        try:
            if self.repo.table_exists('master_ninos'):
                count = self.repo.fetch_one('SELECT COUNT(*) AS total FROM master_ninos') or {'total': 0}
                if int(count.get('total') or 0) > 0:
                    where = ['activo = 1']
                    params = []
                    if uds and uds != 'TODAS':
                        where.append('unidad_servicio = ?')
                        params.append(uds)
                    params.append(limit)
                    return self.repo.fetch_all(
                        f"""
                        SELECT documento, nombre_completo, unidad_servicio AS unidad, coordinador, grupo_etario,
                               estado, peso, talla, perimetro_braquial, diagnostico_nutricional, vacunas,
                               carne_salud, control_crecimiento, fecha_carga AS fecha_valoracion
                        FROM master_ninos
                        WHERE {' AND '.join(where)}
                        ORDER BY unidad_servicio, nombre_completo
                        LIMIT ?
                        """,
                        params,
                    )
        except Exception:
            pass
        return []

    def subir_evidencia(self, entregable_id: int, file, meta: dict[str, Any]) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        nombre_original = file.filename or 'evidencia'
        ext = os.path.splitext(nombre_original.lower())[1]
        if ext not in ALLOWED_EVIDENCE_EXT:
            raise ValueError('Tipo de archivo no permitido para evidencia.')
        actividad = meta.get('actividad') or ent.get('codigo') or 'ENTREGABLE'
        fecha = meta.get('fecha') or datetime.now().strftime('%Y-%m-%d')
        uds = meta.get('uds') or ent.get('uds') or 'UDS'
        consecutivo = (self.repo.fetch_one('SELECT COUNT(*) AS total FROM sn_entregables_evidencias WHERE entregable_id = ?', (entregable_id,)) or {}).get('total') or 0
        nombre_guardado = f"{_slug(actividad)}_{fecha}_{_slug(uds)}_{int(consecutivo)+1:02d}{ext}"
        carpeta = self.upload_folder / f"{ent.get('anio')}_{int(ent.get('mes') or 0):02d}" / _slug(uds) / _slug(ent.get('codigo'))
        carpeta.mkdir(parents=True, exist_ok=True)
        ruta = carpeta / nombre_guardado
        file.save(ruta)
        evidencia_id = self.repo.execute(
            """
            INSERT INTO sn_entregables_evidencias
            (entregable_id, fundacion_id, nombre_original, nombre_guardado, ruta_archivo, tipo, actividad, fecha_actividad,
             uds, responsable, observaciones, fecha_carga, usuario_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (entregable_id, int(current_tenant_id(1) or 1), nombre_original, nombre_guardado, str(ruta), 'evidencia', actividad, fecha, uds,
             meta.get('responsable') or ent.get('responsable'), meta.get('observaciones') or '', now_iso(), meta.get('usuario_id')),
        )
        if not evidencia_id:
            latest = self.repo.fetch_one('SELECT id FROM sn_entregables_evidencias WHERE entregable_id = ? ORDER BY id DESC LIMIT 1', (entregable_id,)) or {}
            evidencia_id = int(latest.get('id') or 0)
        self._touch(entregable_id, 'en proceso')
        return {'id': evidencia_id, 'nombre_guardado': nombre_guardado}

    def validar(self, entregable_id: int, usuario_id: int | None = None) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        pendientes = []
        archivos = ent.get('archivos') or []
        evidencias = ent.get('evidencias') or []
        actividades = ent.get('actividades') or []
        confirmadas = [a for a in actividades if int(a.get('confirmado') or 0) == 1]
        tipos = {a.get('tipo') for a in archivos}
        if (int(ent.get('requiere_acta') or 0) or int(ent.get('requiere_listado') or 0)) and not confirmadas:
            pendientes.append('Falta registrar y confirmar la actividad realizada.')
        if int(ent.get('requiere_acta') or 0) and 'acta' not in tipos:
            pendientes.append('Falta acta generada.')
        if int(ent.get('requiere_listado') or 0) and 'listado' not in tipos:
            pendientes.append('Falta listado de asistencia.')
        if int(ent.get('requiere_oficio') or 0) and 'oficio' not in tipos:
            pendientes.append('Falta oficio de articulación/canalización.')
        if int(ent.get('requiere_formato_excel') or 0) and 'formato' not in tipos:
            pendientes.append('Falta formato Excel asociado.')
        if int(ent.get('requiere_fotos') or 0):
            minimo = int(ent.get('minimo_fotos') or 4)
            if len(evidencias) < minimo:
                pendientes.append(f'Faltan evidencias fotográficas: {len(evidencias)}/{minimo}.')
        valido = not pendientes
        estado = 'completo' if valido else 'pendiente'
        resultado = {
            'valido': valido, 'pendientes': pendientes, 'evidencias': len(evidencias),
            'archivos': len(archivos), 'actividades_confirmadas': len(confirmadas),
        }
        self.repo.execute(
            'INSERT INTO sn_entregables_validaciones (entregable_id, fundacion_id, valido, pendientes_json, resultado_json, fecha_validacion, usuario_id) VALUES (?, ?, ?, ?, ?, ?, ?)',
            (entregable_id, int(current_tenant_id(1) or 1), 1 if valido else 0, json.dumps(pendientes, ensure_ascii=False), json.dumps(resultado, ensure_ascii=False), now_iso(), usuario_id),
        )
        self._touch(entregable_id, estado, 100 if valido else None)
        return resultado

    def generar_acta(self, entregable_id: int, usuario_id: int | None = None) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        usuarios = self.obtener_usuarios_base(ent.get('uds'), limit=80)
        actividades = [a for a in (ent.get('actividades') or []) if int(a.get('confirmado') or 0) == 1]
        if not actividades:
            raise ValueError('Registre y confirme al menos una actividad antes de generar el acta.')
        plantilla = self._plantilla_activa(str(ent.get('codigo') or ''), 'acta')
        path = self._document_path(ent, 'ACTA', '.docx')
        doc = Document(str(plantilla['ruta_archivo']))
        actividad = actividades[0]
        self._rellenar_acta_oficial(doc, ent, actividad)
        self._add_evidence_annex(doc, ent.get('evidencias') or [])
        doc.save(path)
        return self._register_file(entregable_id, 'acta', path, usuario_id, {
            'usuarios': len(usuarios), 'actividades': len(actividades),
            'plantilla_oficial_id': plantilla['id'], 'plantilla_version': plantilla['version'],
        })

        doc = Document()  # pragma: no cover - legado retenido temporalmente
        doc.add_heading('ACTA DE REUNIONES', 0)
        self._add_doc_context(doc, ent)
        doc.add_heading(ent.get('nombre') or '', level=1)
        doc.add_paragraph(f"Actividad: {ent.get('actividad') or ''}")
        for index, actividad in enumerate(actividades, start=1):
            if len(actividades) > 1:
                doc.add_heading(f'Actividad confirmada {index}', level=2)
            horario = ' - '.join(filter(None, [str(actividad.get('hora_inicio') or ''), str(actividad.get('hora_final') or '')]))
            doc.add_paragraph(f"Fecha: {actividad.get('fecha_actividad') or ''}")
            if horario:
                doc.add_paragraph(f"Horario: {horario}")
            doc.add_paragraph(f"Lugar/UCA: {actividad.get('lugar') or ent.get('uds') or ''}")
            doc.add_paragraph(f"Responsable: {actividad.get('responsable') or ent.get('responsable') or ''}")
            if actividad.get('dirigido_a'):
                doc.add_paragraph(f"Dirigido a: {actividad.get('dirigido_a')}")
            doc.add_heading('Objetivo e intencionalidad', level=2)
            doc.add_paragraph(str(actividad.get('objetivo') or ''))
            agenda = self._json_list(actividad.get('agenda_json'))
            if agenda:
                doc.add_heading('Agenda', level=2)
                for item in agenda:
                    doc.add_paragraph(str(item), style='List Number')
            doc.add_heading('Desarrollo de la actividad', level=2)
            doc.add_paragraph(str(actividad.get('desarrollo') or ''))
            if actividad.get('resultados'):
                doc.add_heading('Resultados', level=2)
                doc.add_paragraph(str(actividad.get('resultados')))
            if actividad.get('dificultades'):
                doc.add_heading('Dificultades identificadas', level=2)
                doc.add_paragraph(str(actividad.get('dificultades')))
            if actividad.get('acciones_mejora'):
                doc.add_heading('Acciones de mejora', level=2)
                doc.add_paragraph(str(actividad.get('acciones_mejora')))
        if usuarios:
            doc.add_heading('Participantes / usuarios relacionados', level=2)
            table = doc.add_table(rows=1, cols=4)
            hdr = table.rows[0].cells
            hdr[0].text = 'Documento'; hdr[1].text = 'Nombre'; hdr[2].text = 'UDS/UCA'; hdr[3].text = 'Observación'
            for u in usuarios[:40]:
                row = table.add_row().cells
                row[0].text = str(u.get('documento') or '')
                row[1].text = str(u.get('nombre_completo') or '')
                row[2].text = str(u.get('unidad') or ent.get('uds') or '')
                row[3].text = str(u.get('diagnostico_nutricional') or u.get('estado') or '')
        doc.add_heading('Compromisos', level=2)
        compromisos = [str(a.get('compromisos') or '').strip() for a in actividades if str(a.get('compromisos') or '').strip()]
        for item in compromisos:
            doc.add_paragraph(item, style='List Bullet')
        if not compromisos:
            doc.add_paragraph('Sin compromisos adicionales registrados.')
        self._add_evidence_annex(doc, ent.get('evidencias') or [])
        doc.add_paragraph('\nFirma responsable: ________________________________')
        doc.save(path)
        return self._register_file(entregable_id, 'acta', path, usuario_id, {'usuarios': len(usuarios), 'actividades': len(actividades)})

    def generar_listado(self, entregable_id: int, usuario_id: int | None = None) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        usuarios = self.obtener_usuarios_base(ent.get('uds'), limit=500)
        plantilla = self._plantilla_activa(str(ent.get('codigo') or ''), 'listado')
        path = self._document_path(ent, 'LISTADO_ASISTENCIA', '.xlsx')
        wb = load_workbook(str(plantilla['ruta_archivo']))
        filas = self._rellenar_excel_oficial(wb, ent, usuarios)
        wb.save(path)
        return self._register_file(entregable_id, 'listado', path, usuario_id, {
            'usuarios': len(usuarios), 'plantilla_oficial_id': plantilla['id'],
            'plantilla_version': plantilla['version'], 'filas_diligenciadas': filas,
        })

        wb = Workbook()  # pragma: no cover - legado retenido temporalmente
        ws = wb.active
        ws.title = 'Listado asistencia'
        rows = [
            ['Actividad', ent.get('nombre') or ''],
            ['UDS/UCA', ent.get('uds') or ''],
            ['Mes/Año', f"{ent.get('mes')}/{ent.get('anio')}"],
            [],
            ['No.', 'Documento', 'Nombre completo', 'UDS/UCA', 'Rol', 'Teléfono', 'Firma', 'Observaciones']
        ]
        for row in rows:
            ws.append(row)
        for idx, u in enumerate(usuarios, start=1):
            ws.append([idx, u.get('documento') or '', u.get('nombre_completo') or '', u.get('unidad') or ent.get('uds') or '', 'Participante/Familia', '', '', ''])
        for col in range(1, 9):
            ws.column_dimensions[chr(64+col)].width = 22
        wb.save(path)
        return self._register_file(entregable_id, 'listado', path, usuario_id, {'usuarios': len(usuarios)})

    def generar_oficio(self, entregable_id: int, usuario_id: int | None = None) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        usuarios = self._usuarios_priorizados(ent.get('uds'))
        path = self._document_path(ent, 'OFICIO_CANALIZACION', '.docx')
        doc = Document()
        doc.add_heading('OFICIO DE ARTICULACIÓN / CANALIZACIÓN EN SALUD', 0)
        self._add_doc_context(doc, ent)
        doc.add_paragraph('Asunto: Relación de usuarios para gestión, articulación o seguimiento en salud.')
        doc.add_paragraph('Cordial saludo. De acuerdo con la revisión del componente Salud y Nutrición, se relacionan los usuarios que requieren verificación, orientación o canalización según la información disponible en la Base Maestra.')
        table = doc.add_table(rows=1, cols=5)
        hdr = table.rows[0].cells
        hdr[0].text = 'Documento'; hdr[1].text = 'Nombre'; hdr[2].text = 'UDS/UCA'; hdr[3].text = 'Situación'; hdr[4].text = 'Acción requerida'
        for u in usuarios[:100]:
            row = table.add_row().cells
            row[0].text = str(u.get('documento') or '')
            row[1].text = str(u.get('nombre_completo') or '')
            row[2].text = str(u.get('unidad') or ent.get('uds') or '')
            row[3].text = str(u.get('motivo') or 'Verificación de datos de salud')
            row[4].text = 'Realizar seguimiento y reportar soporte.'
        if not usuarios:
            doc.add_paragraph('No se identificaron usuarios con alertas automáticas para los filtros seleccionados. Este oficio queda como borrador editable.')
        doc.add_paragraph('\nFirma responsable: ________________________________')
        doc.save(path)
        return self._register_file(entregable_id, 'oficio', path, usuario_id, {'usuarios_priorizados': len(usuarios)})

    def generar_formato(self, entregable_id: int, usuario_id: int | None = None) -> dict[str, Any]:
        ent = self.detalle(entregable_id)
        if not ent:
            raise ValueError('Entregable no encontrado.')
        usuarios = self.obtener_usuarios_base(ent.get('uds'), limit=5000)
        plantilla = self._plantilla_activa(str(ent.get('codigo') or ''), 'formato')
        path = self._document_path(ent, 'FORMATO', '.xlsx')
        wb = load_workbook(str(plantilla['ruta_archivo']))
        filas = self._rellenar_excel_oficial(wb, ent, usuarios)
        wb.save(path)
        return self._register_file(entregable_id, 'formato', path, usuario_id, {
            'usuarios': len(usuarios), 'plantilla_oficial_id': plantilla['id'],
            'plantilla_version': plantilla['version'], 'filas_diligenciadas': filas,
        })

        wb = Workbook()  # pragma: no cover - legado retenido temporalmente
        ws = wb.active
        ws.title = 'Formato entregable'
        ws.append(['Entregable', ent.get('nombre') or ''])
        ws.append(['UDS/UCA', ent.get('uds') or ''])
        ws.append(['Mes/Año', f"{ent.get('mes')}/{ent.get('anio')}"])
        ws.append([])
        ws.append(['No.', 'Documento', 'Nombre completo', 'UDS/UCA', 'Grupo etario', 'Peso', 'Talla', 'Perímetro braquial', 'Diagnóstico', 'Vacunas', 'Carné salud', 'Observaciones'])
        for idx, u in enumerate(usuarios, start=1):
            ws.append([
                idx, u.get('documento') or '', u.get('nombre_completo') or '', u.get('unidad') or ent.get('uds') or '',
                u.get('grupo_etario') or '', u.get('peso') or '', u.get('talla') or '', u.get('perimetro_braquial') or '',
                u.get('diagnostico_nutricional') or '', u.get('vacunas') or '', u.get('carne_salud') or '', ''
            ])
        for col in range(1, 13):
            ws.column_dimensions[chr(64+col) if col <= 26 else 'A'].width = 20
        wb.save(path)
        return self._register_file(entregable_id, 'formato', path, usuario_id, {'usuarios': len(usuarios)})

    def generar_matriz(self, filtros: dict[str, Any], usuario_id: int | None = None) -> dict[str, Any]:
        data = self.listar(filtros)
        rows = data['entregables']
        mes = int(filtros.get('mes') or datetime.now().month)
        anio = int(filtros.get('anio') or datetime.now().year)
        path = self.output_folder / f"MATRIZ_CONTROL_ENTREGABLES_SALUD_NUTRICION_{_month_name(mes)}_{anio}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = 'Matriz control'
        ws.append(['No.', 'Entregable', 'UDS/UCA', 'Responsable', 'Evidencia requerida', 'Evidencia cargada', 'Mínimo fotos', 'Fotos cargadas', 'Estado', 'Fecha', 'Observaciones', 'Archivos generados', 'Validación'])
        for idx, r in enumerate(rows, start=1):
            minimo = int(r.get('minimo_fotos') or 0) if int(r.get('requiere_fotos') or 0) else 0
            fotos = int(r.get('fotos_cargadas') or 0)
            ws.append([idx, r.get('nombre'), r.get('uds'), r.get('responsable'), r.get('evidencias_requeridas'), 'Sí' if fotos else 'No', minimo, fotos, r.get('estado'), r.get('fecha_actualizacion') or r.get('fecha_creacion'), r.get('observaciones'), r.get('archivos_generados'), 'OK' if r.get('estado') == 'completo' else 'Pendiente'])
        for col in range(1, 14):
            ws.column_dimensions[chr(64+col)].width = 22
        wb.save(path)
        return self._register_file(None, 'matriz', path, usuario_id, {'total_entregables': len(rows)})

    def generar_informe(self, filtros: dict[str, Any], usuario_id: int | None = None) -> dict[str, Any]:
        data = self.listar(filtros)
        rows = data['entregables']
        resumen = data['resumen']
        mes = int(filtros.get('mes') or datetime.now().month)
        anio = int(filtros.get('anio') or datetime.now().year)
        path = self.output_folder / f"INFORME_ENTREGABLES_SALUD_NUTRICION_{_month_name(mes)}_{anio}.docx"
        doc = Document()
        detalles = [self.detalle(int(r['id'])) or r for r in rows]
        actividades = [
            actividad
            for detalle in detalles
            for actividad in (detalle.get('actividades') or [])
            if int(actividad.get('confirmado') or 0) == 1
        ]
        participantes = sum(int(a.get('participantes_total') or 0) for a in actividades)
        ucas = sorted({str(r.get('uds') or '').strip() for r in rows if str(r.get('uds') or '').strip()})
        doc.add_heading('INFORME MENSUAL', 0)
        doc.add_heading('Componente Salud y Nutrición', level=1)
        doc.add_paragraph(f"Periodo: {_month_name(mes)} {anio}")
        doc.add_paragraph('Estado del documento: borrador generado por la plataforma para revisión y aprobación profesional.')
        doc.add_heading('1. Resumen ejecutivo', level=1)
        doc.add_paragraph(
            f"Durante el periodo se registraron {len(actividades)} actividades confirmadas en "
            f"{len(ucas)} UCA, con {participantes} participaciones reportadas. El tablero presenta "
            f"{resumen.get('completos')} entregables completos y {resumen.get('pendientes')} pendientes, "
            f"para un cumplimiento documental del {resumen.get('porcentaje')}%."
        )
        doc.add_heading('2. Indicadores de cumplimiento', level=1)
        table = doc.add_table(rows=1, cols=5)
        hdr = table.rows[0].cells
        hdr[0].text = 'No.'; hdr[1].text = 'Entregable'; hdr[2].text = 'UDS/UCA'; hdr[3].text = 'Estado'; hdr[4].text = 'Observaciones'
        for idx, r in enumerate(rows, start=1):
            row = table.add_row().cells
            row[0].text = str(idx); row[1].text = str(r.get('nombre') or ''); row[2].text = str(r.get('uds') or ''); row[3].text = str(r.get('estado') or ''); row[4].text = str(r.get('observaciones') or '')
        doc.add_page_break()
        doc.add_heading('3. Desarrollo de actividades confirmadas', level=1)
        chapter = 1
        for detalle in detalles:
            confirmadas = [a for a in (detalle.get('actividades') or []) if int(a.get('confirmado') or 0) == 1]
            if not confirmadas:
                continue
            doc.add_heading(f"3.{chapter} {detalle.get('nombre') or detalle.get('codigo')}", level=2)
            doc.add_paragraph(f"UCA: {detalle.get('uds') or 'TODAS'}")
            for actividad in confirmadas:
                doc.add_heading(str(actividad.get('fecha_actividad') or ''), level=3)
                doc.add_paragraph(f"Responsable: {actividad.get('responsable') or detalle.get('responsable') or ''}")
                doc.add_paragraph(f"Objetivo: {actividad.get('objetivo') or ''}")
                doc.add_paragraph(str(actividad.get('desarrollo') or ''))
                if actividad.get('resultados'):
                    doc.add_paragraph(f"Resultados: {actividad.get('resultados')}")
                if actividad.get('dificultades'):
                    doc.add_paragraph(f"Dificultades: {actividad.get('dificultades')}")
                if actividad.get('acciones_mejora'):
                    doc.add_paragraph(f"Acciones de mejora: {actividad.get('acciones_mejora')}")
                if actividad.get('compromisos'):
                    doc.add_paragraph(f"Compromisos: {actividad.get('compromisos')}")
            self._add_evidence_annex(doc, detalle.get('evidencias') or [], heading='Evidencias del entregable')
            chapter += 1
        if not actividades:
            doc.add_paragraph('No existen actividades confirmadas para el periodo. No se presentan actividades programadas como ejecutadas.')

        doc.add_heading('4. Pendientes y plan de mejora', level=1)
        pending_count = 0
        for detalle in detalles:
            if str(detalle.get('estado') or '').lower() == 'completo':
                continue
            pending_count += 1
            doc.add_paragraph(
                f"{detalle.get('codigo')} - {detalle.get('nombre')} ({detalle.get('uds')}): "
                f"estado {detalle.get('estado') or 'pendiente'}. Evidencia requerida: "
                f"{detalle.get('evidencias_requeridas') or ''}",
                style='List Bullet',
            )
        if not pending_count:
            doc.add_paragraph('No se registran pendientes documentales en el periodo.')
        doc.add_heading('5. Revisión profesional', level=1)
        doc.add_paragraph('Observaciones finales: ______________________________________________________________')
        doc.add_paragraph('\nNombre y firma de la profesional: _________________________________________________')
        doc.add_paragraph('Fecha de aprobación: ____________________')
        doc.save(path)
        return self._register_file(None, 'informe', path, usuario_id, {
            'total_entregables': len(rows), 'actividades_confirmadas': len(actividades),
            'participantes_reportados': participantes, 'estado': 'BORRADOR_PARA_REVISION',
        })

    def generar_zip(self, filtros: dict[str, Any], usuario_id: int | None = None) -> dict[str, Any]:
        data = self.listar(filtros)
        rows = data['entregables']
        mes = int(filtros.get('mes') or datetime.now().month)
        anio = int(filtros.get('anio') or datetime.now().year)
        # Asegurar matriz e informe actualizados.
        matriz = self.generar_matriz(filtros, usuario_id)
        informe = self.generar_informe(filtros, usuario_id)
        zip_path = self.output_folder / f"ENTREGABLES_SALUD_NUTRICION_{_month_name(mes)}_{anio}.zip"
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.write(matriz['ruta_archivo'], arcname=os.path.basename(matriz['ruta_archivo']))
            zf.write(informe['ruta_archivo'], arcname=os.path.basename(informe['ruta_archivo']))
            for r in rows:
                folder = f"{str(r.get('codigo') or '').replace('E','').split('_')[0].zfill(2)}_{_slug(r.get('nombre'))}/"
                for a in self.repo.fetch_all('SELECT * FROM sn_entregables_archivos WHERE entregable_id = ?', (r['id'],)):
                    ruta = a.get('ruta_archivo')
                    if ruta and os.path.exists(ruta):
                        zf.write(ruta, arcname=folder + os.path.basename(ruta))
                for e in self.repo.fetch_all('SELECT * FROM sn_entregables_evidencias WHERE entregable_id = ?', (r['id'],)):
                    ruta = e.get('ruta_archivo')
                    if ruta and os.path.exists(ruta):
                        zf.write(ruta, arcname=folder + 'evidencias/' + os.path.basename(ruta))
        return self._register_file(None, 'zip', zip_path, usuario_id, {'total_entregables': len(rows)})

    def _register_file(self, entregable_id: int | None, tipo: str, path: Path | str, usuario_id: int | None, metadata: dict[str, Any]) -> dict[str, Any]:
        path = Path(path)
        archivo_id = self.repo.execute(
            """
            INSERT INTO sn_entregables_archivos
            (entregable_id, fundacion_id, tipo, nombre_archivo, ruta_archivo, estado, metadata_json, fecha_generacion, usuario_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (entregable_id, int(current_tenant_id(1) or 1), tipo, path.name, str(path), 'generado', json.dumps(metadata, ensure_ascii=False), now_iso(), usuario_id),
        )
        if not archivo_id:
            if entregable_id:
                latest = self.repo.fetch_one('SELECT id FROM sn_entregables_archivos WHERE entregable_id = ? ORDER BY id DESC LIMIT 1', (entregable_id,)) or {}
            else:
                latest = self.repo.fetch_one('SELECT id FROM sn_entregables_archivos WHERE ruta_archivo = ? ORDER BY id DESC LIMIT 1', (str(path),)) or {}
            archivo_id = int(latest.get('id') or 0)
        if entregable_id:
            self._touch(entregable_id, 'en proceso')
        return {'archivo_id': archivo_id, 'tipo': tipo, 'nombre_archivo': path.name, 'ruta_archivo': str(path), 'download_url': f'/api/salud-nutricion/entregables/archivo/{archivo_id}'}

    def _document_path(self, ent: dict[str, Any], prefix: str, ext: str) -> Path:
        folder = self.output_folder / f"{ent.get('anio')}_{int(ent.get('mes') or 0):02d}" / _slug(ent.get('uds')) / _slug(ent.get('codigo'))
        folder.mkdir(parents=True, exist_ok=True)
        name = f"{prefix}_{_slug(ent.get('codigo'))}_{_slug(ent.get('uds'))}_{datetime.now().strftime('%Y%m%d%H%M%S')}{ext}"
        return folder / name

    @staticmethod
    def _json_list(value: Any) -> list[Any]:
        if isinstance(value, list):
            return value
        try:
            parsed = json.loads(value or '[]')
            return parsed if isinstance(parsed, list) else []
        except (TypeError, ValueError, json.JSONDecodeError):
            return []

    @staticmethod
    def _normalizar_etiqueta(value: Any) -> str:
        import unicodedata
        text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
        return re.sub(r'\s+', ' ', text).strip().lower().rstrip(':')

    def _rellenar_acta_oficial(self, doc: Document, ent: dict[str, Any], actividad: dict[str, Any]) -> None:
        """Completa celdas rotuladas sin reconstruir tablas, logos ni estilos oficiales."""
        valores = {
            'fecha': actividad.get('fecha_actividad'),
            'hora inicio': actividad.get('hora_inicio'),
            'hora inicial': actividad.get('hora_inicio'),
            'hora final': actividad.get('hora_final'),
            'lugar': actividad.get('lugar') or ent.get('uds'),
            'nombre del responsable': actividad.get('responsable') or ent.get('responsable'),
            'responsable': actividad.get('responsable') or ent.get('responsable'),
            'tipo de actividad': ent.get('nombre'),
            'actividad': ent.get('nombre'),
            'dirigido a': actividad.get('dirigido_a'),
            'tema': ent.get('nombre'),
            'objetivo': actividad.get('objetivo'),
            'agenda': '\n'.join(str(x) for x in self._json_list(actividad.get('agenda_json'))),
            'desarrollo': actividad.get('desarrollo'),
            'desarrollo de la actividad': actividad.get('desarrollo'),
            'compromisos': actividad.get('compromisos'),
        }
        pendientes = {key: str(value or '').strip() for key, value in valores.items() if str(value or '').strip()}
        for table in doc.tables:
            for row_index, row in enumerate(table.rows):
                for col_index, cell in enumerate(row.cells):
                    etiqueta = self._normalizar_etiqueta(cell.text)
                    clave = next((key for key in pendientes if etiqueta == key or etiqueta.startswith(key + ' ')), None)
                    if not clave:
                        continue
                    candidates = list(row.cells[col_index + 1:])
                    if row_index + 1 < len(table.rows) and col_index < len(table.rows[row_index + 1].cells):
                        candidates.append(table.rows[row_index + 1].cells[col_index])
                    target = next((candidate for candidate in candidates if candidate._tc is not cell._tc), None)
                    if target is not None:
                        target.text = pendientes.pop(clave)

    def _rellenar_excel_oficial(self, wb, ent: dict[str, Any], usuarios: list[dict[str, Any]]) -> int:
        """Diligencia columnas reconocibles sin crear hojas ni alterar el diseño del libro oficial."""
        aliases = {
            'documento': ('documento', 'numero de documento', 'no documento', 'identificacion'),
            'nombre_completo': ('nombre completo', 'nombres y apellidos', 'nombre del participante', 'usuario'),
            'unidad': ('uds', 'uca', 'unidad de servicio', 'nombre de la unidad'),
            'grupo_etario': ('grupo etario', 'grupo de edad'),
            'peso': ('peso', 'peso kg', 'peso actual'),
            'talla': ('talla', 'longitud', 'talla cm'),
            'perimetro_braquial': ('perimetro braquial', 'pb', 'circunferencia braquial'),
            'diagnostico_nutricional': ('diagnostico nutricional', 'clasificacion nutricional', 'estado nutricional'),
            'vacunas': ('vacunas', 'esquema de vacunacion'),
            'carne_salud': ('carne de salud', 'carne crecimiento y desarrollo', 'crecimiento y desarrollo'),
        }
        total = 0
        for ws in wb.worksheets:
            header_row = None
            mapping: dict[int, str] = {}
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
                found: dict[int, str] = {}
                for cell in row:
                    text = self._normalizar_etiqueta(cell.value)
                    for field, names in aliases.items():
                        if any(name == text or name in text for name in names):
                            found[cell.column] = field
                            break
                if len(found) >= 2:
                    header_row, mapping = row[0].row, found
                    break
            if not header_row:
                continue
            for offset, usuario in enumerate(usuarios, start=1):
                target_row = header_row + offset
                for column, field in mapping.items():
                    value = usuario.get(field)
                    if field == 'unidad':
                        value = value or ent.get('uds')
                    cell = ws.cell(target_row, column)
                    if not isinstance(cell, MergedCell):
                        cell.value = value or ''
                total += 1
        return total

    def _add_evidence_annex(self, doc: Document, evidencias: list[dict[str, Any]], heading: str = 'Registro fotográfico') -> None:
        imagenes = []
        for evidencia in evidencias:
            ruta = Path(str(evidencia.get('ruta_archivo') or ''))
            if ruta.suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp'} and ruta.exists():
                imagenes.append((ruta, evidencia))
        if not imagenes:
            return
        doc.add_heading(heading, level=2)
        for ruta, evidencia in imagenes:
            try:
                doc.add_picture(str(ruta), width=Inches(5.8))
                doc.add_paragraph(
                    f"{evidencia.get('actividad') or 'Actividad'} · "
                    f"{evidencia.get('fecha_actividad') or ''} · {evidencia.get('uds') or ''}"
                )
            except Exception:
                doc.add_paragraph(f"Evidencia adjunta: {evidencia.get('nombre_guardado') or ruta.name}")

    def _touch(self, entregable_id: int, estado: str | None = None, porcentaje: int | None = None) -> None:
        if estado is not None and porcentaje is not None:
            self.repo.execute('UPDATE sn_entregables_mes SET estado = ?, porcentaje = ?, fecha_actualizacion = ? WHERE id = ?', (estado, porcentaje, now_iso(), entregable_id))
        elif estado is not None:
            self.repo.execute('UPDATE sn_entregables_mes SET estado = ?, fecha_actualizacion = ? WHERE id = ?', (estado, now_iso(), entregable_id))
        else:
            self.repo.execute('UPDATE sn_entregables_mes SET fecha_actualizacion = ? WHERE id = ?', (now_iso(), entregable_id))

    def _add_doc_context(self, doc: Document, ent: dict[str, Any]) -> None:
        doc.add_paragraph(f"Periodo: {_month_name(int(ent.get('mes') or 1))} {ent.get('anio')}")
        doc.add_paragraph(f"UDS/UCA: {ent.get('uds') or 'TODAS'}")
        doc.add_paragraph(f"Responsable: {ent.get('responsable') or 'Enfermera/Nutricionista'}")
        doc.add_paragraph(f"Fecha de generación: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    def _usuarios_priorizados(self, uds: str | None = None) -> list[dict[str, Any]]:
        usuarios = self.obtener_usuarios_base(uds, limit=1000)
        priorizados = []
        for u in usuarios:
            motivos = []
            for campo, label in [('vacunas', 'Sin soporte de vacunas'), ('carne_salud', 'Sin carné de salud'), ('diagnostico_nutricional', 'Sin diagnóstico nutricional')]:
                val = str(u.get(campo) or '').strip().lower()
                if not val or val in {'no', 'sin dato', 'pendiente', 'none'}:
                    motivos.append(label)
            diag = str(u.get('diagnostico_nutricional') or '').lower()
            if any(t in diag for t in ['desnutric', 'riesgo', 'severa', 'moderada']):
                motivos.append('Alerta nutricional')
            if motivos:
                x = dict(u)
                x['motivo'] = '; '.join(motivos)
                priorizados.append(x)
        return priorizados

    def archivo(self, archivo_id: int) -> dict[str, Any] | None:
        return self.repo.fetch_one('SELECT * FROM sn_entregables_archivos WHERE id = ?', (archivo_id,))
