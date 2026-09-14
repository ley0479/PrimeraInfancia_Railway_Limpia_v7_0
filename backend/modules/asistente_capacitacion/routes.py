from __future__ import annotations

from datetime import datetime, timedelta
from flask import Blueprint, jsonify, request, g, Response
from modules.dbapi_compat import sqlite3
from modules.seguridad.services import ROLE_MENU_PERMISSIONS, get_request_user_context
from .guides import DEFAULT_GUIDE, GUIDES
from .schema import SCHEMA_SQL
from .config import public_flags, public_liam_flags, public_elian_flags
from .elian_module_registry import authorized_modules
from .assistant_service import respond
from .platform_profile import get_platform_profile
from .tool_registry import ALLOWED_TOOLS, MODULE_DATASETS, execute
from .rate_limit import allow
from .provider_adapter import OpenAIResponsesProvider, ProviderUnavailable, provider_status
from .knowledge_base import manual_for_role, manual_for_question, build_manual_pdf
from .privacy_service import redact
from .local_speech import enabled as local_speech_enabled, status as local_speech_status, transcribe_wav
from .action_intents import propose_action, propose_read_actions
from .error_center import record as record_incident, get as get_incident, list_recent as list_incidents
from .credit_agent import parse_credit_request, query as query_credits, create_proposal as create_credit_proposal, confirm as confirm_credit_proposal
from .action_policy import decision as action_decision, public_policy
from .system_prompt import realtime_instructions
from .orchestrator import LiamOrchestrator
import csv, io, json, uuid, os, tempfile, re, hashlib, requests


def register_asistente_capacitacion(app, database_path: str) -> None:
    def connect():
        conn = sqlite3.connect(database_path); conn.row_factory = sqlite3.Row; return conn

    conn = connect(); conn.executescript(SCHEMA_SQL); conn.commit(); conn.close()
    bp = Blueprint('asistente_capacitacion', __name__, url_prefix='/api/asistente-capacitacion')
    orchestrator=LiamOrchestrator(database_path)

    def audit_lia(ctx, event_type, *, module=None, tool=None, success=True, request_id=None, metadata=None):
        conn=None
        try:
            conn=connect();conn.execute('''INSERT INTO lia_audit_events
              (fundacion_id,usuario_id,event_type,modulo,tool_name,success,request_id,metadata_redacted,created_at)
              VALUES(?,?,?,?,?,?,?,?,?)''',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0),event_type,module,tool,1 if success else 0,request_id,json.dumps(metadata or {},ensure_ascii=False),datetime.now().isoformat(timespec='seconds')));conn.commit()
        except Exception as exc:
            if conn:
                try: conn.rollback()
                except Exception: pass
            app.logger.warning('LÍA continúa sin auditoría auxiliar: %s', type(exc).__name__)
        finally:
            if conn:
                try: conn.close()
                except Exception: pass

    def limited(ctx):
        flags=public_flags();key=f"{ctx.get('fundacion_id')}:{ctx.get('usuario_id')}"
        return not allow(key,flags['rate_limit_per_minute'])

    def visual_payload(result, module):
        """Contrato visual cerrado; el cliente nunca interpreta HTML procedente del modelo."""
        tool='';component='spotlight';data={};display='inline'
        def table_for(value):
            if not isinstance(value,dict):return None
            if isinstance(value.get('relation_rows'),list):
                rows=list(value.get('relation_rows') or [])
                if value.get('total_row'):rows.append(value['total_row'])
                return {'title':value.get('title'),'note':value.get('rule'),'metrics':value.get('metrics') or [],'columns':value.get('columns') or [],'rows':rows,'maxColumns':19,'relationFormat':True}
            if {'profiles','beneficiaries','units'} <= set(value):
                p=value.get('profiles') or {};b=value.get('beneficiaries') or {};u=value.get('units') or {};rows=[['Beneficiarios',b.get('total',0),'Base Maestra'],['UDS activas',u.get('registered_active',0),'Unidades'],['Perfiles',p.get('total',0),'Usuarios'],['Coordinadores',p.get('coordinators',0),'Consolidado'],['Talento humano',p.get('interdisciplinary_team_total',0),'Base institucional']]
                rows.extend([[f"UDS: {x.get('unit') or 'Sin nombre'}",x.get('beneficiaries',0),x.get('coordinator') or 'Sin coordinador'] for x in (u.get('items') or [])[:20]]);return {'metrics':[{'label':'Beneficiarios','value':b.get('total',0)},{'label':'UDS activas','value':u.get('registered_active',0)},{'label':'Coordinadores','value':p.get('coordinators',0)},{'label':'Talento humano','value':p.get('interdisciplinary_team_total',0)},{'label':'Perfiles','value':p.get('total',0)},{'label':'Fuentes cargadas','value':(value.get('sources') or {}).get('total_sources',0)},{'label':'Movimientos','value':(value.get('movements') or {}).get('total',0)}],'columns':['Indicador','Valor','Detalle'],'rows':rows}
            if isinstance(value.get('beneficiaries'),list):return {'columns':['Nombre','Documento','UDS','Estado'],'rows':[[x.get('nombre_completo'),x.get('documento'),x.get('unidad_servicio'),x.get('estado')] for x in value['beneficiaries'][:30]]}
            if isinstance(value.get('results'),list):return {'columns':['Tipo','Resultado','Referencia','Unidad/Rol','Estado','Fuente'],'rows':[[x.get('resource_type'),x.get('title'),x.get('reference'),x.get('unit'),x.get('status'),x.get('source')] for x in value['results'][:30]]}
            if isinstance(value.get('comparison'),list):return {'title':f"Comparación {value.get('period_a')} / {value.get('period_b')}",'note':value.get('disclaimer'),'columns':['Indicador',value.get('period_a'),value.get('period_b'),'Variación','Disponibilidad'],'rows':[[x.get('indicator'),x.get('period_a'),x.get('period_b'),x.get('variation'),'Completa' if x.get('available') else 'Sin datos'] for x in value['comparison']]}
            if isinstance(value.get('rows'),list) and isinstance(value.get('fields'),list):return {'title':value.get('title') or 'Vista previa del reporte','note':f"Fuente: {value.get('source')}. Vista previa de solo lectura.",'columns':[str(x).replace('_',' ').title() for x in value['fields']],'rows':[[item.get(field) for field in value['fields']] for item in value['rows']]}
            if isinstance(value.get('summary'),dict) and isinstance(value.get('items'),list) and value.get('source')=='Calendario Inteligente':
                s=value['summary'];return {'title':'Supervisión de entregables','metrics':[{'label':'Esperados','value':s.get('expected',0)},{'label':'Recibidos','value':s.get('received',0)},{'label':'Pendientes','value':s.get('pending',0)},{'label':'Vencidos','value':s.get('overdue',0)},{'label':'Devueltos','value':s.get('returned',0)},{'label':'Incompletos','value':s.get('incomplete',0)},{'label':'Cumplimiento','value':f"{s.get('compliance_percent',0)}%"}],'columns':['Entregable','Unidad','Responsable','Fecha límite','Clasificación','Prioridad'],'rows':[[x.get('title'),x.get('unit'),x.get('responsible'),x.get('due_date'),x.get('category'),x.get('priority')] for x in value['items']]}
            if isinstance(value.get('components'),list) and value.get('sanitized') is True:return {'title':'Salud del sistema','metrics':[{'label':'Estado general','value':value.get('overall_status')},{'label':'Eventos fallidos','value':value.get('failed_audit_events') if value.get('failed_audit_events') is not None else 'No disponible'}],'columns':['Componente','Estado'],'rows':[[x.get('component'),x.get('status')] for x in value['components']]}
            if isinstance(value.get('summary'),dict) and isinstance(value.get('items'),list) and value.get('source')=='Fundaciones y Suscripciones':
                s=value['summary'];return {'title':'Fundaciones y licencias','metrics':[{'label':'Fundaciones','value':s.get('total',0)},{'label':'Activas','value':s.get('active',0)},{'label':'Próximas a vencer','value':s.get('expiring',0)},{'label':'Vencidas','value':s.get('expired',0)},{'label':'Sin usuarios','value':s.get('without_users',0)},{'label':'Sin actividad','value':s.get('without_activity',0)},{'label':'Crédito bajo','value':s.get('low_credit',0)}],'columns':['Fundación','Estado','Suscripción','Plan','Vencimiento','Días','Créditos','Usuarios','Alerta'],'rows':[[x.get('foundation'),x.get('status'),x.get('subscription_status'),x.get('plan'),x.get('expiry_date'),x.get('days_remaining'),x.get('credits_available'),x.get('active_users'),x.get('alert')] for x in value['items']]}
            if isinstance(value.get('findings'),list) and value.get('source')=='Base Maestra activa':return {'title':'Calidad de Base Maestra','metrics':[{'label':'Registros analizados','value':value.get('total_records',0)},{'label':'Alertas','value':(value.get('summary') or {}).get('alerts',0)},{'label':'Duplicados','value':(value.get('summary') or {}).get('duplicate_document_groups',0)},{'label':'UDS inexistentes','value':(value.get('summary') or {}).get('unregistered_units',0)}],'columns':['Hallazgo','Total','Severidad'],'rows':[[x.get('label'),x.get('total'),x.get('severity')] for x in value['findings']]}
            if isinstance(value.get('warnings'),list) and value.get('predictive_language')=='risk_only':return {'title':'Alertas tempranas','note':value.get('disclaimer'),'metrics':[{'label':'Alertas','value':(value.get('summary') or {}).get('total_warnings',0)},{'label':'Críticas','value':(value.get('summary') or {}).get('critico',0)},{'label':'Altas','value':(value.get('summary') or {}).get('alto',0)},{'label':'Preventivas','value':(value.get('summary') or {}).get('preventivo',0)}],'columns':['Nivel','Riesgo','Registros','Fuente','Motivo'],'rows':[[x.get('level'),x.get('title'),x.get('total'),x.get('source'),x.get('reason')] for x in value['warnings']]}
            if isinstance(value.get('incidents'),list) and value.get('sanitized') is True:return {'title':'Centro de incidencias','metrics':[{'label':'Incidencias','value':(value.get('summary') or {}).get('total',0)},{'label':'Abiertas','value':(value.get('summary') or {}).get('open',0)},{'label':'En análisis','value':(value.get('summary') or {}).get('in_analysis',0)},{'label':'Resueltas','value':(value.get('summary') or {}).get('resolved',0)}],'columns':['Incidencia','Módulo','Código','Tipo','Estado','Severidad','Fecha'],'rows':[[x.get('incident_id'),x.get('module'),x.get('error_code'),x.get('error_type'),x.get('status'),x.get('severity'),x.get('created_at')] for x in value['incidents']]}
            if isinstance(value.get('notifications'),list) and value.get('send_actions') is False:return {'title':'Centro de notificaciones','metrics':[{'label':'Pendientes','value':(value.get('summary') or {}).get('total',0)},{'label':'Críticas','value':(value.get('summary') or {}).get('critical',0)},{'label':'Advertencias','value':(value.get('summary') or {}).get('warning',0)},{'label':'Información','value':(value.get('summary') or {}).get('information',0)}],'columns':['Prioridad','Origen','Notificación','Estado','Fecha'],'rows':[[x.get('priority'),x.get('source'),x.get('title'),x.get('status'),x.get('date')] for x in value['notifications']]}
            if isinstance(value.get('profiles'),list):return {'columns':['Nombre','Usuario','Rol','Estado'],'rows':[[x.get('nombre_completo'),x.get('username'),x.get('rol'),'Activo' if x.get('activo') else 'Inactivo'] for x in value['profiles'][:30]]}
            if isinstance(value.get('indicators'),dict):
                metrics=[{'label':str(k).replace('_',' ').title(),'value':v} for k,v in value['indicators'].items() if isinstance(v,(int,float))]
                return {'metrics':metrics[:16],'columns':['Indicador','Total'],'rows':[[x['label'],x['value']] for x in metrics]}
            if isinstance(value.get('datasets'),list):return {'columns':['Fuente','Registros','Disponibilidad'],'rows':[[x.get('name'),x.get('total'), 'Disponible' if x.get('available',True) else 'No disponible'] for x in value['datasets'][:30]]}
            if isinstance(value.get('items'),list):return {'columns':['Actividad','Fecha','Estado','Unidad'],'rows':[[x.get('title'),x.get('due_date'),x.get('status'),x.get('unit')] for x in value['items'][:30]]}
            return None
        if result.get('tool_results'):
            combined=[]
            for entry in result['tool_results'][:8]:
                current=table_for(entry.get('result'));tool=str(entry.get('tool') or tool)
                if current:
                    for row in current['rows'][:20]:combined.append([tool,*row])
            if combined:component='table';display='drawer';data={'columns':['Consulta','Dato 1','Dato 2','Dato 3','Dato 4'],'rows':combined[:50]}
        if not data:
            value=result.get('tool_result') if isinstance(result.get('tool_result'),dict) else {};table=table_for(value)
            if table:component='metric-card' if table.get('metrics') else 'table';display='drawer';data=table
        if not data:
            steps=[x.strip(' -') for x in re.split(r'\n+|(?=\d+\.\s)',str(result.get('message') or '')) if x.strip()]
            if len(steps)>1:component='list';data={'items':[{'label':f'Paso {i}','value':text} for i,text in enumerate(steps[:12],1)]}
        target=f'#nav-{module}' if re.fullmatch(r'[a-z0-9-]+',str(module or '')) else None
        return {'text':result.get('message') or '','componentType':component,'data':data,'targetSelector':target,'display':display,'schemaVersion':'lia-ui-v1','supportedComponents':['metric-card','table','list','spotlight'],'tool':tool or None}

    def save_message(ctx, *, role, content, module, request_id):
        if role not in {'user','assistant'}:raise ValueError('Rol de conversación no válido.')
        conn=connect()
        try:
            conn.execute('INSERT INTO lia_conversation_messages(fundacion_id,usuario_id,role,content_redacted,module,request_id,created_at) VALUES(?,?,?,?,?,?,?)',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0),role,redact(content)[:10000],module,request_id,datetime.now().isoformat(timespec='seconds')))
            conn.commit()
        except Exception:conn.rollback();raise
        finally:conn.close()

    def save_exchange(ctx, *, question, answer, module, request_id):
        """Conserva cada turno como filas independientes; nunca reemplaza turnos previos."""
        now=datetime.now().isoformat(timespec='seconds');conn=connect()
        try:
            values=(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0),module,request_id,now)
            conn.execute('INSERT INTO lia_conversation_messages(fundacion_id,usuario_id,role,content_redacted,module,request_id,created_at) VALUES(?,?,?,?,?,?,?)',(values[0],values[1],'user',redact(question)[:5000],*values[2:]))
            conn.execute('INSERT INTO lia_conversation_messages(fundacion_id,usuario_id,role,content_redacted,module,request_id,created_at) VALUES(?,?,?,?,?,?,?)',(values[0],values[1],'assistant',redact(answer)[:10000],*values[2:]))
            conn.commit()
        except Exception:
            conn.rollback();raise
        finally:conn.close()

    @bp.get('/chat/history')
    def chat_history():
        ctx=get_request_user_context()
        try:limit=max(1,min(int(request.args.get('limit') or 100),500))
        except (TypeError,ValueError):limit=100
        try:page=max(1,int(request.args.get('page') or 1))
        except (TypeError,ValueError):page=1
        search=str(request.args.get('search') or '').strip()[:120];module_filter=str(request.args.get('module') or '').strip()[:80];role_filter=str(request.args.get('role') or '').strip().lower();date_from=str(request.args.get('date_from') or '').strip()[:10];date_to=str(request.args.get('date_to') or '').strip()[:10]
        if role_filter and role_filter not in {'user','assistant'}:return jsonify({'error':'Tipo de mensaje no válido.'}),422
        include_archived=str(request.args.get('include_archived') or '').lower() in {'1','true','yes'}
        where='fundacion_id=? AND usuario_id=?';params=[int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0)]
        if not include_archived:where+=' AND NOT EXISTS (SELECT 1 FROM lia_conversation_archives a WHERE a.fundacion_id=lia_conversation_messages.fundacion_id AND a.usuario_id=lia_conversation_messages.usuario_id AND a.request_id=lia_conversation_messages.request_id)'
        if search:where+=' AND LOWER(content_redacted) LIKE LOWER(?)';params.append(f'%{search}%')
        if module_filter:where+=' AND module=?';params.append(module_filter)
        if role_filter:where+=' AND role=?';params.append(role_filter)
        if date_from:where+=' AND created_at>=?';params.append(date_from+'T00:00:00')
        if date_to:where+=' AND created_at<=?';params.append(date_to+'T23:59:59')
        conn=connect();total=int(conn.execute(f'SELECT COUNT(*) FROM lia_conversation_messages WHERE {where}',tuple(params)).fetchone()[0] or 0)
        rows=conn.execute(f'''SELECT id,role,content_redacted,module,request_id,created_at
          FROM lia_conversation_messages WHERE {where}
          ORDER BY id DESC LIMIT ? OFFSET ?''',tuple([*params,limit,(page-1)*limit])).fetchall();conn.close()
        messages=[{'id':row['id'],'role':row['role'],'content':row['content_redacted'],'module':row['module'],'request_id':row['request_id'],'created_at':row['created_at']} for row in reversed(rows)]
        return jsonify({'messages':messages,'total':total,'page':page,'limit':limit,'has_more':page*limit<total,'filters':{'search':search or None,'module':module_filter or None,'role':role_filter or None,'date_from':date_from or None,'date_to':date_to or None},'append_only':True}),200

    @bp.get('/chat/history/export.csv')
    def export_chat_history():
        ctx=get_request_user_context();conn=connect();rows=conn.execute('''SELECT created_at,role,module,content_redacted FROM lia_conversation_messages
          WHERE fundacion_id=? AND usuario_id=? ORDER BY id''',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0))).fetchall();conn.close()
        stream=io.StringIO();writer=csv.writer(stream);writer.writerow(['fecha','rol','modulo','mensaje'])
        for row in rows:writer.writerow([row['created_at'],row['role'],row['module'],row['content_redacted']])
        return Response('\ufeff'+stream.getvalue(),200,{'Content-Type':'text/csv; charset=utf-8','Content-Disposition':'attachment; filename="historial_lia.csv"','Cache-Control':'no-store'})

    @bp.get('/chat/history/export.xlsx')
    def export_chat_history_xlsx():
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        ctx=get_request_user_context();conn=connect();rows=conn.execute('''SELECT created_at,role,module,content_redacted FROM lia_conversation_messages
          WHERE fundacion_id=? AND usuario_id=? ORDER BY id''',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0))).fetchall();conn.close()
        book=Workbook();sheet=book.active;sheet.title='Historial Lía';sheet.append(['Fecha','Tipo','Módulo','Mensaje'])
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='0F766E')
        for row in rows:sheet.append([row['created_at'],'Usuario' if row['role']=='user' else 'Lía',row['module'],row['content_redacted']])
        sheet.column_dimensions['A'].width=22;sheet.column_dimensions['B'].width=12;sheet.column_dimensions['C'].width=28;sheet.column_dimensions['D'].width=90
        output=io.BytesIO();book.save(output);output.seek(0)
        return Response(output.getvalue(),200,{'Content-Type':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','Content-Disposition':'attachment; filename="historial_lia.xlsx"','Cache-Control':'no-store'})

    @bp.get('/chat/history/export.pdf')
    def export_chat_history_pdf():
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        ctx=get_request_user_context();conn=connect();rows=conn.execute('''SELECT created_at,role,module,content_redacted FROM lia_conversation_messages
          WHERE fundacion_id=? AND usuario_id=? ORDER BY id''',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0))).fetchall();conn.close()
        output=io.BytesIO();doc=SimpleDocTemplate(output,pagesize=letter,title='Historial de Lía');styles=getSampleStyleSheet();story=[Paragraph('Historial de conversaciones con Lía',styles['Title']),Spacer(1,12)]
        for row in rows:
            heading=f"{row['created_at']} · {'Orden' if row['role']=='user' else 'Respuesta'} · {row['module'] or 'sin módulo'}"
            story.extend([Paragraph(heading,styles['Heading4']),Paragraph(str(row['content_redacted']).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'),styles['BodyText']),Spacer(1,8)])
        doc.build(story)
        return Response(output.getvalue(),200,{'Content-Type':'application/pdf','Content-Disposition':'attachment; filename="historial_lia.pdf"','Cache-Control':'no-store'})

    @bp.get('/chat/history/sessions')
    def conversation_sessions():
        ctx=get_request_user_context();include_archived=str(request.args.get('include_archived') or '').lower() in {'1','true','yes'};fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0);conn=connect()
        rows=conn.execute('''SELECT m.request_id,MIN(m.created_at) AS started_at,MAX(m.created_at) AS ended_at,COUNT(*) AS messages,
          SUM(CASE WHEN m.role='user' THEN 1 ELSE 0 END) AS commands,MAX(m.module) AS module,MAX(CASE WHEN a.request_id IS NULL THEN 0 ELSE 1 END) AS archived
          FROM lia_conversation_messages m LEFT JOIN lia_conversation_archives a ON a.fundacion_id=m.fundacion_id AND a.usuario_id=m.usuario_id AND a.request_id=m.request_id
          WHERE m.fundacion_id=? AND m.usuario_id=? GROUP BY m.request_id ORDER BY MAX(m.id) DESC LIMIT 200''',(fid,uid)).fetchall();conn.close()
        sessions=[dict(row) for row in rows if include_archived or not row['archived']]
        return jsonify({'sessions':sessions,'total':len(sessions),'include_archived':include_archived}),200

    @bp.post('/chat/history/sessions/<string:request_id>/archive')
    def archive_conversation(request_id):
        ctx=get_request_user_context();fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0);conn=connect();exists=conn.execute('SELECT 1 FROM lia_conversation_messages WHERE fundacion_id=? AND usuario_id=? AND request_id=?',(fid,uid,request_id)).fetchone()
        if not exists:conn.close();return jsonify({'error':'Conversación no encontrada.'}),404
        now=datetime.now().isoformat(timespec='seconds');conn.execute('INSERT INTO lia_conversation_archives(fundacion_id,usuario_id,request_id,archived_at) VALUES(?,?,?,?) ON CONFLICT(fundacion_id,usuario_id,request_id) DO UPDATE SET archived_at=excluded.archived_at',(fid,uid,request_id,now));conn.commit();conn.close()
        return jsonify({'message':'Conversación archivada sin eliminar sus mensajes.','request_id':request_id,'archived':True}),200

    @bp.delete('/chat/history/sessions/<string:request_id>/archive')
    def restore_conversation(request_id):
        ctx=get_request_user_context();conn=connect();conn.execute('DELETE FROM lia_conversation_archives WHERE fundacion_id=? AND usuario_id=? AND request_id=?',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0),request_id));conn.commit();conn.close()
        return jsonify({'message':'Conversación restaurada.','request_id':request_id,'archived':False}),200

    @bp.get('/chat/history/stats')
    def conversation_stats():
        ctx=get_request_user_context();fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0);conn=connect()
        totals=conn.execute('''SELECT COUNT(*) AS messages,SUM(CASE WHEN role='user' THEN 1 ELSE 0 END) AS commands,COUNT(DISTINCT request_id) AS sessions FROM lia_conversation_messages WHERE fundacion_id=? AND usuario_id=?''',(fid,uid)).fetchone()
        modules=conn.execute('''SELECT COALESCE(NULLIF(module,''),'sin módulo') AS module,COUNT(*) AS total FROM lia_conversation_messages WHERE fundacion_id=? AND usuario_id=? GROUP BY module ORDER BY total DESC LIMIT 10''',(fid,uid)).fetchall();conn.close()
        return jsonify({'messages':int(totals['messages'] or 0),'commands':int(totals['commands'] or 0),'sessions':int(totals['sessions'] or 0),'top_modules':[dict(row) for row in modules],'scope':{'foundation_id':fid,'user_id':uid}}),200

    @bp.get('/chat/history/admin')
    def admin_chat_history():
        ctx=get_request_user_context()
        if str(ctx.get('rol') or '').upper() not in {'SUPERADMIN','GERENTE','COORDINADOR'}:return jsonify({'error':'No tienes permiso para auditar conversaciones.'}),403
        try:limit=max(1,min(int(request.args.get('limit') or 100),500))
        except (TypeError,ValueError):limit=100
        target=request.args.get('user_id',type=int);where='m.fundacion_id=?';params=[int(ctx.get('fundacion_id') or 1)]
        if target:where+=' AND m.usuario_id=?';params.append(target)
        conn=connect();rows=conn.execute(f'''SELECT m.id,m.usuario_id,u.username,m.role,m.content_redacted,m.module,m.request_id,m.created_at
          FROM lia_conversation_messages m LEFT JOIN usuarios_app u ON u.id=m.usuario_id AND u.fundacion_id=m.fundacion_id
          WHERE {where} ORDER BY m.id DESC LIMIT ?''',tuple([*params,limit])).fetchall();conn.close()
        audit_lia(ctx,'CONVERSATION_HISTORY_AUDITED',module='administracion',metadata={'target_user_id':target,'rows':len(rows)})
        return jsonify({'messages':[dict(row) for row in rows],'scope':{'foundation_id':int(ctx.get('fundacion_id') or 1),'cross_foundation':False},'read_only':True}),200

    @bp.get('/config')
    def config_publica():
        return jsonify({'lia': public_flags(), 'liam': public_liam_flags(), 'elian': public_elian_flags(), 'platform_profile': get_platform_profile()}), 200

    @bp.get('/contexto')
    def contexto():
        if not public_flags()['enabled']:
            return jsonify({'error':'LÍA está desactivada.'}), 404
        ctx = get_request_user_context(); modulo = str(request.args.get('modulo') or 'dashboard').strip()
        allowed = set(ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''), []))
        if allowed and modulo not in allowed: return jsonify({'error':'Módulo no autorizado para el rol actual.'}), 403
        guide = dict(GUIDES.get(modulo, DEFAULT_GUIDE)); guide['modulo'] = modulo
        guide['tour_steps']=[{'help_id':f'{modulo}.screen','message':guide.get('resumen') or ''}]+[{'help_id':f'{modulo}.primary-action','message':step} for step in guide.get('pasos',[])]
        knowledge = manual_for_role(str(ctx.get('rol') or ''), module_id=modulo, screen_id=str(request.args.get('screen_id') or ''), help_id=str(request.args.get('help_id') or ''))
        conn = connect(); row = conn.execute('SELECT * FROM ayuda_progreso_usuario WHERE fundacion_id=? AND usuario_id=? AND modulo=?',(ctx.get('fundacion_id') or 1,ctx.get('usuario_id'),modulo)).fetchone(); conn.close()
        return jsonify({'guia':guide,'conocimiento':knowledge,'rol':ctx.get('rol'),'progreso':dict(row) if row else None,'solo_orientacion':True}), 200

    @bp.get('/presentation')
    def presentation():
        if not public_flags()['enabled']: return jsonify({'error':'LÍA está desactivada.'}),404
        ctx=get_request_user_context();allowed=list(ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''),[]))
        modules=[]
        for key in allowed:
            item=GUIDES.get(key)
            if item: modules.append({'module':key,'title':item.get('titulo'),'purpose':item.get('proposito') or item.get('resumen')})
        profile=get_platform_profile();audit_lia(ctx,'PLATFORM_PRESENTATION_OPENED',module='dashboard',metadata={'modules':len(modules)})
        workflow=['Confirmar sesión, fundación y periodo.','Actualizar las fuentes autorizadas en Base Maestra.','Revisar unidades, participantes y equipo humano.','Consultar calendario, actividades y entregables.','Trabajar en el módulo correspondiente según el rol.','Cargar evidencias o generar borradores.','Confirmar el resultado y atender revisiones o devoluciones.']
        return jsonify({'profile':profile,'modules':modules,'role':ctx.get('rol'),'total':len(modules),'workflow':workflow}),200

    @bp.get('/manual')
    def manual_maestro():
        ctx = get_request_user_context()
        manual = manual_for_role(
            str(ctx.get('rol') or ''),
            module_id=str(request.args.get('module_id') or '').strip(),
            screen_id=str(request.args.get('screen_id') or '').strip(),
            help_id=str(request.args.get('help_id') or '').strip(),
        )
        audit_lia(ctx, 'MASTER_MANUAL_OPENED', module=request.args.get('module_id') or 'manual-operativo', metadata={'help_id': request.args.get('help_id') or ''})
        return jsonify(manual), 200

    @bp.get('/manual.pdf')
    def manual_maestro_pdf():
        ctx = get_request_user_context()
        manual = manual_for_role(str(ctx.get('rol') or ''))
        payload = build_manual_pdf(manual)
        audit_lia(ctx, 'MASTER_MANUAL_PDF_DOWNLOADED', module='manual-operativo', metadata={'bytes': len(payload)})
        return Response(payload, 200, {
            'Content-Type': 'application/pdf',
            'Content-Disposition': 'attachment; filename="Manual_Maestro_Primera_Infancia.pdf"',
            'Cache-Control': 'no-store',
        })

    @bp.get('/elian/platform-tour')
    def elian_platform_tour():
        if not public_elian_flags()['enabled'] or not public_elian_flags()['platform_tour_enabled']:
            return jsonify({'error':'El recorrido general de ELIAN está desactivado.'}),404
        ctx=get_request_user_context();fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0)
        allowed=ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''),[])
        modules=authorized_modules(allowed)
        conn=connect();row=conn.execute('SELECT * FROM elian_platform_tour_progress WHERE fundacion_id=? AND usuario_id=? AND tour_id=?',(fid,uid,'platform-overview')).fetchone();conn.close()
        progress=dict(row) if row else None
        if progress:
            for field in ('completed_modules_json','skipped_modules_json','pending_modules_json'):
                progress[field[:-5]]=json.loads(progress.pop(field) or '[]')
        audit_lia(ctx,'ELIAN_PLATFORM_TOUR_OPENED',module='dashboard',metadata={'modules':len(modules)})
        return jsonify({'tour_id':'platform-overview','tour_version':1,'profile':get_platform_profile(),'role':ctx.get('rol'),'modules':modules,'total':len(modules),'progress':progress,'navigation_policy':'registered_routes_only'}),200

    @bp.route('/elian/platform-tour/progress',methods=['GET','PUT'])
    def elian_platform_tour_progress():
        if not public_elian_flags()['enabled']: return jsonify({'error':'ELIAN está desactivado.'}),404
        ctx=get_request_user_context();fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0);tour_id='platform-overview';conn=connect()
        if request.method=='GET':
            row=conn.execute('SELECT * FROM elian_platform_tour_progress WHERE fundacion_id=? AND usuario_id=? AND tour_id=?',(fid,uid,tour_id)).fetchone();conn.close()
            if not row:return jsonify({'progress':None}),200
            value=dict(row)
            for field in ('completed_modules_json','skipped_modules_json','pending_modules_json'):value[field[:-5]]=json.loads(value.pop(field) or '[]')
            return jsonify({'progress':value}),200
        data=request.get_json(silent=True) or {};allowed_ids=[m['module_id'] for m in authorized_modules(ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''),[]))];allowed_set=set(allowed_ids)
        status=str(data.get('status') or 'in_progress');mode=str(data.get('mode') or 'automatic')
        if status not in {'not_started','in_progress','paused','completed','cancelled','outdated','failed'} or mode not in {'automatic','interactive','pending','module'}:
            conn.close();return jsonify({'error':'Estado o modo de recorrido no válido.'}),422
        current=str(data.get('current_module_id') or '')
        if current and current not in allowed_set:conn.close();return jsonify({'error':'Módulo no autorizado para el recorrido.'}),403
        def clean_list(name):
            values=data.get(name) or []
            return [item for item in dict.fromkeys(str(v) for v in values) if item in allowed_set]
        completed=clean_list('completed_modules');skipped=clean_list('skipped_modules');pending=[m for m in allowed_ids if m not in set(completed+skipped)]
        now=datetime.now().isoformat(timespec='seconds');completed_at=now if status=='completed' else None
        conn.execute('''INSERT INTO elian_platform_tour_progress(fundacion_id,usuario_id,tour_id,tour_version,current_module_id,current_step,completed_modules_json,skipped_modules_json,pending_modules_json,mode,status,created_at,updated_at,completed_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(fundacion_id,usuario_id,tour_id) DO UPDATE SET tour_version=excluded.tour_version,current_module_id=excluded.current_module_id,current_step=excluded.current_step,completed_modules_json=excluded.completed_modules_json,skipped_modules_json=excluded.skipped_modules_json,pending_modules_json=excluded.pending_modules_json,mode=excluded.mode,status=excluded.status,updated_at=excluded.updated_at,completed_at=excluded.completed_at''',(fid,uid,tour_id,1,current,max(0,int(data.get('current_step') or 0)),json.dumps(completed),json.dumps(skipped),json.dumps(pending),mode,status,now,now,completed_at));conn.commit();conn.close()
        audit_lia(ctx,'ELIAN_PLATFORM_TOUR_PROGRESS',module=current or 'dashboard',metadata={'status':status,'completed':len(completed),'skipped':len(skipped)})
        return jsonify({'message':'Progreso de ELIAN actualizado.','progress':{'tour_id':tour_id,'tour_version':1,'current_module_id':current,'current_step':max(0,int(data.get('current_step') or 0)),'completed_modules':completed,'skipped_modules':skipped,'pending_modules':pending,'mode':mode,'status':status}}),200

    @bp.route('/elian/visual-config',methods=['GET','PUT'])
    def elian_visual_config():
        ctx=get_request_user_context();fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0)
        variants={
            'afro_colombian_institutional':{'label':'Afrocolombiano institucional','assets':{'male':'./assets/lia/elian-afro-institutional-male-v1.png','female':'./assets/lia/liam-afro-institutional-fullbody-v2.png'},'ready_genders':['male','female']},
            'afro_colombian_technological':{'label':'Afrocolombiano tecnológico','assets':{'male':'./assets/lia/elian-afro-technological-male-v1.png','female':'./assets/lia/elian-afro-technological-female-v1.png'},'ready_genders':['male','female']},
            'afro_colombian_educational':{'label':'Afrocolombiano educativo','assets':{'male':'./assets/lia/elian-afro-educational-male-v1.png','female':'./assets/lia/elian-afro-educational-female-v1.png'},'ready_genders':['male','female']},
        }
        defaults={'assistant_name':'LIAM','avatar_gender':'female','avatar_variant':'afro_colombian_institutional','skin_tone':'dark','hair_style':'short_coily','clothing_style':'institutional_vest','primary_color':'#123A63','secondary_color':'#16C6D8','voice_gender':'female','voice_speed':.95,'headset_enabled':1,'tablet_enabled':1,'hologram_enabled':1,'animation_enabled':1,'walk_enabled':1,'lip_sync_enabled':1,'motion_level':'full','avatar_asset_path':variants['afro_colombian_institutional']['assets']['female']}
        conn=connect();row=conn.execute('SELECT * FROM elian_visual_configuration WHERE fundacion_id=?',(fid,)).fetchone()
        if request.method=='GET':
            conn.close();config={**defaults,**(dict(row) if row else {})}
            if str(config.get('assistant_name') or '').upper() in {'IAN','ELIAN'}:
                config.update({'assistant_name':'LIAM','avatar_gender':'female','voice_gender':'female','motion_level':'full','walk_enabled':1,'avatar_asset_path':variants['afro_colombian_institutional']['assets']['female']})
            if config.get('avatar_gender')=='female' and config.get('avatar_variant')=='afro_colombian_institutional':
                config['avatar_asset_path']=variants['afro_colombian_institutional']['assets']['female']
            selected=variants.get(config['avatar_variant'],variants['afro_colombian_institutional']);ready=config['avatar_gender'] in selected['ready_genders'];return jsonify({'configuration':config,'variants':variants,'genders':['male','female'],'editable':str(ctx.get('rol') or '') in {'SUPERADMIN','GERENTE'},'fallback_active':not ready}),200
        if str(ctx.get('rol') or '') not in {'SUPERADMIN','GERENTE'}:conn.close();return jsonify({'error':'Solo un administrador autorizado puede cambiar la apariencia global.'}),403
        data=request.get_json(silent=True) or {};gender=str(data.get('avatar_gender') or defaults['avatar_gender']);variant=str(data.get('avatar_variant') or defaults['avatar_variant']);motion=str(data.get('motion_level') or defaults['motion_level'])
        if gender not in {'male','female'} or variant not in variants or motion not in {'full','light','reduced'}:conn.close();return jsonify({'error':'La variante visual solicitada no está registrada.'}),422
        asset=variants[variant]['assets'][gender];asset_ready=gender in variants[variant]['ready_genders'];name=str(data.get('assistant_name') or 'LIAM').strip()[:40] or 'LIAM';now=datetime.now().isoformat(timespec='seconds')
        values=(name,gender,variant,str(data.get('skin_tone') or 'dark')[:30],str(data.get('hair_style') or 'short_coily')[:40],str(data.get('clothing_style') or 'institutional_vest')[:40],str(data.get('primary_color') or '#123A63')[:16],str(data.get('secondary_color') or '#16C6D8')[:16],str(data.get('voice_gender') or gender)[:12],max(.6,min(1.5,float(data.get('voice_speed') or .95))),1 if data.get('headset_enabled',True) else 0,1 if data.get('tablet_enabled',True) else 0,1 if data.get('hologram_enabled',True) else 0,1 if data.get('animation_enabled',True) else 0,1 if data.get('walk_enabled',False) else 0,1 if data.get('lip_sync_enabled',False) else 0,motion,asset)
        conn.execute('''INSERT INTO elian_visual_configuration(fundacion_id,assistant_name,avatar_gender,avatar_variant,skin_tone,hair_style,clothing_style,primary_color,secondary_color,voice_gender,voice_speed,headset_enabled,tablet_enabled,hologram_enabled,animation_enabled,walk_enabled,lip_sync_enabled,motion_level,avatar_asset_path,updated_by,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(fundacion_id) DO UPDATE SET assistant_name=excluded.assistant_name,avatar_gender=excluded.avatar_gender,avatar_variant=excluded.avatar_variant,skin_tone=excluded.skin_tone,hair_style=excluded.hair_style,clothing_style=excluded.clothing_style,primary_color=excluded.primary_color,secondary_color=excluded.secondary_color,voice_gender=excluded.voice_gender,voice_speed=excluded.voice_speed,headset_enabled=excluded.headset_enabled,tablet_enabled=excluded.tablet_enabled,hologram_enabled=excluded.hologram_enabled,animation_enabled=excluded.animation_enabled,walk_enabled=excluded.walk_enabled,lip_sync_enabled=excluded.lip_sync_enabled,motion_level=excluded.motion_level,avatar_asset_path=excluded.avatar_asset_path,updated_by=excluded.updated_by,updated_at=excluded.updated_at''',(fid,*values,uid,now,now));conn.commit();conn.close();audit_lia(ctx,'ELIAN_VISUAL_CONFIGURATION_UPDATED',module='administracion',metadata={'variant':variant,'gender':gender,'asset_ready':asset_ready})
        saved=dict(zip(('assistant_name','avatar_gender','avatar_variant','skin_tone','hair_style','clothing_style','primary_color','secondary_color','voice_gender','voice_speed','headset_enabled','tablet_enabled','hologram_enabled','animation_enabled','walk_enabled','lip_sync_enabled','motion_level','avatar_asset_path'),values))
        return jsonify({'message':'Configuración visual de LIAM actualizada.','configuration':saved,'asset_ready':asset_ready}),200

    @bp.post('/progreso')
    def progreso():
        if not public_flags()['enabled']:
            return jsonify({'error':'LÍA está desactivada.'}), 404
        ctx=get_request_user_context(); data=request.get_json(silent=True) or {}; modulo=str(data.get('modulo') or '').strip()
        if not modulo: return jsonify({'error':'Módulo requerido.'}),400
        allowed=set(ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''), []))
        if allowed and modulo not in allowed: return jsonify({'error':'Módulo no autorizado.'}),403
        now=datetime.now().isoformat(timespec='seconds'); completed=1 if data.get('recorrido_completado') else 0; skipped=1 if data.get('recorrido_omitido') else 0
        conn=connect(); conn.execute("""INSERT INTO ayuda_progreso_usuario
        (fundacion_id,usuario_id,modulo,recorrido_completado,recorrido_omitido,veces_abierto,ultima_apertura,fecha_creacion,fecha_actualizacion)
        VALUES (?,?,?,?,?,1,?,?,?) ON CONFLICT(fundacion_id,usuario_id,modulo) DO UPDATE SET
        recorrido_completado=CASE WHEN excluded.recorrido_completado=1 THEN 1 ELSE ayuda_progreso_usuario.recorrido_completado END,
        recorrido_omitido=CASE WHEN excluded.recorrido_omitido=1 THEN 1 ELSE ayuda_progreso_usuario.recorrido_omitido END,
        veces_abierto=ayuda_progreso_usuario.veces_abierto+1,ultima_apertura=excluded.ultima_apertura,fecha_actualizacion=excluded.fecha_actualizacion""",
        (ctx.get('fundacion_id') or 1,ctx.get('usuario_id'),modulo,completed,skipped,now,now,now)); conn.commit(); conn.close()
        return jsonify({'message':'Progreso de aprendizaje actualizado.'}),200

    @bp.post('/chat')
    def chat():
        flags = public_flags()
        if not flags['enabled'] or not flags['text_enabled']:
            return jsonify({'error':'El chat de LÍA está desactivado.'}), 404
        ctx = get_request_user_context()
        if limited(ctx): return jsonify({'error':'Demasiadas solicitudes a LÍA. Espera un momento.'}),429
        data = request.get_json(silent=True) or {}
        question = str(data.get('message') or '').strip()
        module = str(data.get('module') or 'dashboard').strip()
        if not question:
            return jsonify({'error':'La pregunta es obligatoria.'}), 400
        if len(question) > flags['max_message_length']:
            return jsonify({'error':'La pregunta supera el límite permitido.'}), 413
        allowed = set(ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''), []))
        if allowed and module not in allowed:
            return jsonify({'error':'Módulo no autorizado para el rol actual.'}), 403
        knowledge=manual_for_question(str(ctx.get('rol') or ''),module_id=module,screen_id=str(data.get('screen_id') or ''),help_id=str(data.get('help_id') or ''))
        history=[]
        for item in (data.get('history') or [])[-6:]:
            if not isinstance(item,dict) or item.get('role') not in {'user','assistant'}: continue
            content=redact(str(item.get('content') or '').strip())[:1200]
            if content: history.append({'role':item['role'],'content':content})
        screen_context=data.get('screen_context') if isinstance(data.get('screen_context'),dict) else {}
        proposal=propose_action(question,screen_context=screen_context)
        read_plan=propose_read_actions(question,screen_context=screen_context)
        credit_request=parse_credit_request(question)
        result=respond(question=question, module=module, role=str(ctx.get('rol') or ''),allowed_modules=sorted(allowed),knowledge=knowledge,history=history)
        if len(read_plan)>1 and not credit_request:
            user=dict(getattr(g,'current_user',None) or {}) or {'id':ctx.get('usuario_id'),'rol':ctx.get('rol')}
            tool_results=[];parts=[]
            for step in read_plan[:5]:
                try:
                    value=execute(step['server_tool'],args=step.get('arguments') or {},database_path=database_path,tenant_id=int(ctx.get('fundacion_id') or 1),user=user)
                    tool_results.append({'tool':step['server_tool'],'result':value})
                    if step['server_tool']=='get_foundation_data_summary':
                        p=value['profiles'];b=value['beneficiaries'];u=value['units']
                        coord_names=', '.join(x.get('name') or '' for x in p.get('coordinator_items') or []) or 'sin nombres registrados'
                        unit_names=', '.join(x.get('unit') or '' for x in u.get('items') or []) or 'sin nombres registrados'
                        check=value.get('consistency') or {};warning=' '.join(check.get('warnings') or [])
                        parts.append(f"Base Maestra: {b['total']} beneficiarios, {u['registered_active']} UDS activas ({unit_names}), {p['total']} perfiles, {p['coordinators']} coordinadores ({coord_names}) y {p.get('interdisciplinary_team_total',0)} integrantes de talento humano. Consistencia: {check.get('status','sin validar')}"+(f"; {warning}" if warning else ''))
                    elif step['server_tool']=='list_foundation_profiles':parts.append(f"Perfiles encontrados: {value['total']}")
                    elif step['server_tool']=='search_foundation_beneficiaries':
                        names=', '.join(str(x.get('nombre_completo') or x.get('documento')) for x in value.get('beneficiaries',[])[:5]) or 'sin coincidencias'
                        parts.append(f"Beneficiarios encontrados: {value['total']} ({names})")
                    elif step['server_tool']=='universal_search':parts.append(f"Búsqueda universal: {value['total']} resultados autorizados para {value.get('query')}")
                    elif step['server_tool']=='get_pending_activities_summary':parts.append(f"Tareas pendientes: {value['total']}, de ellas {value['overdue']} vencidas y {value['due_today']} para hoy")
                    elif step['server_tool']=='get_platform_module_summary':
                        detail=', '.join(f"{item['name']}: {item['total']}" for item in value.get('datasets') or [])
                        parts.append(f"{value['module']}: {detail}")
                    elif step['server_tool']=='get_monthly_health_indicators':
                        x=value['indicators'];parts.append(f"Salud: {x['carne_salud']} con carné, {x['crecimiento_desarrollo']} con crecimiento y desarrollo, {x['registro_civil']} con registro civil, {x['perimetro_braquial']} con perímetro braquial, {x['gestantes_control_prenatal']} gestantes con control prenatal, {x['sobrepeso']} con sobrepeso, {x['desnutricion']} con desnutrición y {x['riesgo_desnutricion']} en riesgo")
                    elif step['server_tool']=='get_monthly_relation_summary':
                        metrics={x['label']:x['value'] for x in value.get('metrics') or []};parts.append(f"Relación del Mes {value.get('period')}: {metrics.get('Unidades de atención',0)} unidades, {metrics.get('Total usuarios',0)} usuarios, {metrics.get('Total huevos',0)} huevos, {metrics.get('Cubetas de 30',0)} cubetas, {metrics.get('Panales completos',0)} panales y {metrics.get('Total verduras',0)} verduras")
                    audit_lia(ctx,'TOOL_COMPLETED',module=module,tool=step['server_tool'],request_id=result['request_id'],metadata={'read_only':True,'agentic_step':len(tool_results)})
                except (PermissionError,LookupError,ValueError) as exc:parts.append(f"{step['server_tool']}: {exc}")
            message='. '.join(parts)+'.'
            result.update({'message':message,'speech_text':message,'confidence':'confirmed','confirmation_required':False,'actions':[],'tool_results':tool_results,'agentic':{'steps':len(tool_results),'max_steps':5,'read_only':True}})
            result['ui']=visual_payload(result,module)
            result.update(result['ui'])
            audit_lia(ctx,'QUESTION_COMPLETED',module=module,request_id=result['request_id'],metadata={'length':len(question),'multi_tool':True,'steps':len(tool_results)})
            save_exchange(ctx,question=question,answer=result['message'],module=module,request_id=result['request_id'])
            return jsonify(result),200
        incident_match=re.search(r'\bINC-\d{8}-\d{6}-[A-Z0-9]{6}\b',question.upper())
        if incident_match:
            incident=get_incident(database_path,incident_match.group(0),int(ctx.get('fundacion_id') or 1))
            if incident:
                diagnostic_message=f"El incidente {incident['incident_id']} corresponde a {incident['error_type'].replace('_',' ')}. Causa: {incident['cause']} Solución: {incident['solution']}"
                result.update({'message':diagnostic_message,'speech_text':diagnostic_message,'diagnostic':incident,'confidence':'confirmed','confirmation_required':False,'actions':[]})
            else:
                result.update({'message':'No encontré ese incidente dentro de tu fundación o no tienes autorización para consultarlo.','speech_text':'No encontré ese incidente dentro de tu fundación o no tienes autorización para consultarlo.','confidence':'insufficient','confirmation_required':False,'actions':[]})
            proposal=None
        if credit_request:
            try:
                if credit_request['kind']=='query':
                    credit_result=query_credits(database_path,credit_request,str(ctx.get('rol') or ''),int(ctx.get('fundacion_id') or 1));message=credit_result['message'];result.update({'message':message,'speech_text':message,'tool_result':credit_result,'confidence':'confirmed','confirmation_required':False,'actions':[]});audit_lia(ctx,'TOOL_COMPLETED',module='facturacion',tool='consultar_creditos',request_id=result['request_id'],metadata={'scope':credit_result['scope']})
                elif str(ctx.get('rol') or '')!='SUPERADMIN':
                    result.update({'message':'Solo SUPERADMIN puede modificar créditos o suscripciones mediante Lian.','speech_text':'Solo SUPERADMIN puede modificar créditos o suscripciones mediante Lian.','confidence':'forbidden','confirmation_required':False,'actions':[]})
                else:
                    credit_proposal=create_credit_proposal(database_path,credit_request,int(ctx.get('usuario_id') or 0),str(ctx.get('rol') or ''));result.update({'message':credit_proposal['summary']+' ¿Confirmas?','speech_text':credit_proposal['summary']+' ¿Confirmas?','confidence':'confirmed','confirmation_required':True,'action_proposal':credit_proposal,'actions':[]})
            except (PermissionError,LookupError,ValueError) as exc:
                result.update({'message':str(exc),'speech_text':str(exc),'confidence':'needs_input','confirmation_required':False,'actions':[]})
            proposal=None
        if proposal:
            policy=action_decision(proposal.get('id'),str(ctx.get('rol') or ''))
            if not policy.get('allowed'):
                proposal=None
                result.update({'message':policy.get('reason'),'speech_text':policy.get('reason'),'confidence':'forbidden','confirmation_required':False,'actions':[]})
            else:
                proposal.update({'risk':policy['risk'],'confirmation_type':policy['confirmation']})
                if proposal.get('confirmation_required'):
                    proposal.setdefault('expires_at',(datetime.now()+timedelta(seconds=60)).isoformat(timespec='seconds'))
        if proposal:
            target_module=str((proposal.get('arguments') or {}).get('module') or '')
            if target_module and allowed and target_module not in allowed:
                proposal=None
                result.update({'message':'Tu rol no tiene permiso para abrir o consultar ese módulo.','speech_text':'Tu rol no tiene permiso para abrir o consultar ese módulo.','confidence':'forbidden','confirmation_required':False,'actions':[]})
            elif proposal.get('server_tool'):
                try:
                    user=dict(getattr(g,'current_user',None) or {}) or {'id':ctx.get('usuario_id'),'rol':ctx.get('rol')}
                    tool_result=execute(proposal['server_tool'],args=proposal.get('arguments') or {},database_path=database_path,tenant_id=int(ctx.get('fundacion_id') or 1),user=user)
                    if proposal['server_tool']=='get_foundation_data_summary':
                        profiles=tool_result.get('profiles') or {};beneficiaries=tool_result.get('beneficiaries') or {};units=tool_result.get('units') or {}
                        groups=', '.join(f"{item['age_group']}: {item['total']}" for item in beneficiaries.get('by_age_group') or []) or 'sin grupos registrados'
                        coord_names=', '.join(x.get('name') or '' for x in profiles.get('coordinator_items') or []) or 'sin nombres registrados'
                        unit_names=', '.join(x.get('unit') or '' for x in units.get('items') or []) or 'sin nombres registrados'
                        sources=', '.join(f"{x.get('source')}: {x.get('valid',0)} válidos" for x in (tool_result.get('sources') or {}).get('active_loads') or []) or 'sin cargas registradas'
                        moves=', '.join(f"{x.get('type')}: {x.get('total',0)}" for x in (tool_result.get('movements') or {}).get('by_type') or []) or 'sin movimientos'
                        check=tool_result.get('consistency') or {};warning=' '.join(check.get('warnings') or [])
                        message=f"Resumen de la fundación activa: {profiles.get('total',0)} perfiles; {profiles.get('coordinators',0)} coordinadores ({coord_names}); {profiles.get('interdisciplinary_team_total',0)} integrantes de talento humano; {beneficiaries.get('total',0)} beneficiarios; y {units.get('registered_active',0)} UDS ({unit_names}). Por grupo etario: {groups}. Cargas vigentes: {sources}. Movimientos de la versión maestra: {moves}. Consistencia: {check.get('status','sin validar')}."+(f" Advertencia: {warning}" if warning else '')
                        actions=[];total=int(beneficiaries.get('total') or 0)
                    elif proposal['server_tool']=='list_foundation_profiles':
                        roles={}
                        for item in tool_result.get('profiles') or []:roles[item.get('rol') or 'SIN_ROL']=roles.get(item.get('rol') or 'SIN_ROL',0)+1
                        detail=', '.join(f'{role}: {count}' for role,count in sorted(roles.items())) or 'sin perfiles'
                        message=f"Encontré {tool_result.get('total',0)} perfiles en la fundación de tu sesión. En esta página: {detail}."
                        actions=[];total=int(tool_result.get('total') or 0)
                    elif proposal['server_tool']=='search_foundation_beneficiaries':
                        total=int(tool_result.get('total') or 0);items=tool_result.get('beneficiaries') or []
                        detail=', '.join(f"{item.get('nombre_completo') or 'Sin nombre'} ({item.get('documento') or 'sin documento'}, {item.get('unidad_servicio') or 'sin UDS'})" for item in items[:10]) or 'sin coincidencias'
                        message=f"Encontré {total} beneficiarios autorizados: {detail}."
                        actions=[]
                    elif proposal['server_tool']=='universal_search':
                        total=int(tool_result.get('total') or 0);detail=', '.join(f"{x.get('title')} ({x.get('resource_type')})" for x in tool_result.get('results') or []) or 'sin coincidencias'
                        message=f"Encontré {total} resultados autorizados en tu fundación: {detail}. Fuente indicada en la tabla."
                        actions=[]
                    elif proposal['server_tool']=='get_platform_module_summary':
                        total=int(tool_result.get('total_records') or 0)
                        detail=', '.join(f"{item['name']}: {item['total']}" for item in tool_result.get('datasets') or []) or 'sin registros'
                        message=f"Resumen de {tool_result.get('module')}: {detail}."
                        actions=[]
                    elif proposal['server_tool']=='get_monthly_health_indicators':
                        x=tool_result['indicators'];total=int(x.get('total') or 0)
                        message=f"De {total} beneficiarios: {x['carne_salud']} tienen carné de salud, {x['crecimiento_desarrollo']} crecimiento y desarrollo, {x['registro_civil']} registro civil, {x['perimetro_braquial']} perímetro braquial y {x['gestantes_control_prenatal']} de {x['gestantes']} gestantes tienen control prenatal. Anexo nutricional: {x['sobrepeso']} con sobrepeso, {x['desnutricion']} con desnutrición y {x['riesgo_desnutricion']} en riesgo."
                        actions=[]
                    elif proposal['server_tool']=='get_monthly_relation_summary':
                        values={x['label']:x['value'] for x in tool_result.get('metrics') or []};total=int(values.get('Total usuarios') or 0)
                        message=f"Relación del Mes {tool_result.get('period')}: {values.get('Unidades de atención',0)} unidades de atención y {total} usuarios. La entrega corresponde a {values.get('Total huevos',0)} huevos, {values.get('Cubetas de 30',0)} cubetas, {values.get('Panales completos',0)} panales completos, {values.get('Total verduras',0)} verduras, {values.get('Ollas comunitarias',0)} ollas comunitarias y {values.get('Bienestarina',0)} unidades de Bienestarina. Te muestro el detalle por UDS con el mismo orden del formato oficial."
                        actions=[]
                    elif proposal['server_tool']=='compare_periods':
                        total=len(tool_result.get('comparison') or []);available=sum(1 for x in tool_result.get('comparison') or [] if x.get('available'))
                        message=f"Comparación {tool_result.get('period_a')} contra {tool_result.get('period_b')}: {available} de {total} indicadores tienen datos en ambos periodos. Los faltantes no fueron estimados."
                        actions=[]
                    elif proposal['server_tool']=='build_custom_report_preview':
                        message=f"Preparé la vista previa {tool_result.get('title')}: {tool_result.get('shown_rows',0)} de {tool_result.get('total_rows',0)} filas, desde {tool_result.get('source')}. No se modificó ni exportó información."
                        actions=[]
                    elif proposal['server_tool']=='supervise_deliverables':
                        summary=tool_result.get('summary') or {};message=f"Supervisión de entregables: {summary.get('received',0)} recibidos de {summary.get('expected',0)} esperados; {summary.get('pending',0)} pendientes, {summary.get('overdue',0)} vencidos, {summary.get('returned',0)} devueltos y {summary.get('incomplete',0)} incompletos. Cumplimiento: {summary.get('compliance_percent',0)}%."
                        actions=[]
                    elif proposal['server_tool']=='get_system_health':
                        message=f"Estado general del sistema: {tool_result.get('overall_status')}. Revisé {len(tool_result.get('components') or [])} componentes con un diagnóstico sanitizado que no contiene secretos."
                        actions=[]
                    elif proposal['server_tool']=='get_foundation_portfolio':
                        summary=tool_result.get('summary') or {};message=f"Portafolio de fundaciones: {summary.get('total',0)} registradas, {summary.get('active',0)} activas, {summary.get('expiring',0)} próximas a vencer, {summary.get('expired',0)} vencidas, {summary.get('without_users',0)} sin usuarios y {summary.get('without_activity',0)} sin actividad."
                        actions=[]
                    elif proposal['server_tool']=='analyze_master_data_quality':
                        summary=tool_result.get('summary') or {};message=f"Calidad de Base Maestra: analicé {tool_result.get('total_records',0)} registros y encontré {summary.get('alerts',0)} tipos de alerta, {summary.get('duplicate_document_groups',0)} grupos de documentos duplicados y {summary.get('unregistered_units',0)} beneficiarios con UDS inexistente. No realicé correcciones automáticas."
                        actions=[]
                    elif proposal['server_tool']=='get_early_warnings':
                        summary=tool_result.get('summary') or {};message=f"Según los datos disponibles, se detectan {summary.get('total_warnings',0)} señales de riesgo: {summary.get('critico',0)} críticas, {summary.get('alto',0)} altas y {summary.get('preventivo',0)} preventivas. No son predicciones ni diagnósticos."
                        actions=[]
                    elif proposal['server_tool']=='get_incident_center':
                        summary=tool_result.get('summary') or {};message=f"Centro de incidencias: {summary.get('total',0)} resultados dentro de tu alcance; {summary.get('open',0)} abiertos, {summary.get('in_analysis',0)} en análisis, {summary.get('in_progress',0)} en proceso y {summary.get('resolved',0)} resueltos."
                        actions=[]
                    elif proposal['server_tool']=='get_notification_center':
                        summary=tool_result.get('summary') or {};message=f"Centro de notificaciones: {summary.get('total',0)} avisos pendientes; {summary.get('critical',0)} críticos, {summary.get('warning',0)} advertencias y {summary.get('information',0)} informativos. Solo los consulté; no envié ni marqué ninguno."
                        actions=[]
                    else:
                        total=int(tool_result.get('total') or 0); overdue=int(tool_result.get('overdue') or 0); today=int(tool_result.get('due_today') or 0); upcoming=int(tool_result.get('upcoming') or 0); undated=int(tool_result.get('undated') or 0)
                        query=tool_result.get('query') or {}; scope_label='del equipo' if query.get('scope')=='team' else 'asignadas a tu cuenta'; period_label=f" del periodo {query.get('period')}" if query.get('period') else ''
                        message=f'Encontré {total} actividades {scope_label}{period_label}: {overdue} vencidas, {today} para hoy, {upcoming} próximas y {undated} sin fecha.'
                        actions=[{'type':'navigate','module':'calendario-inteligente','period':query.get('period'),'scope':query.get('scope'),'target':'calendario.pending.list'}]
                    result.update({'message':message,'speech_text':message,'confidence':'confirmed','confirmation_required':False,'tool_result':tool_result,'actions':actions})
                    audit_lia(ctx,'TOOL_COMPLETED',module=module,tool=proposal['server_tool'],request_id=result['request_id'],metadata={'read_only':True,'total':total})
                except (PermissionError,LookupError,ValueError) as exc:
                    result.update({'message':str(exc),'speech_text':str(exc),'confidence':'insufficient','confirmation_required':False})
            elif proposal['missing']:
                result.update({'message':proposal['summary']+' Antes de hacerlo necesito: '+', '.join(proposal['missing'])+'.','speech_text':proposal['summary']+' Antes de hacerlo necesito '+', '.join(proposal['missing'])+'.','confidence':'needs_input','confirmation_required':False,'action_proposal':proposal})
            elif not proposal.get('confirmation_required'):
                result.update({'message':proposal['summary'],'speech_text':proposal['summary'],'confidence':'confirmed','confirmation_required':False,'actions':[{'type':proposal['client_handler'],**proposal['arguments']}]})
            else:
                result.update({'message':proposal['summary']+' Revisa los datos y confirma para continuar.','speech_text':proposal['summary']+' Revisa los datos y confirma para continuar.','confirmation_required':True,'action_proposal':proposal})
        if not proposal and flags['ai_enabled'] and provider_status()['ready']:
            try:
                generated=OpenAIResponsesProvider().respond(messages=[*history,{'role':'user','content':redact(question)}],context={
                    'role':str(ctx.get('rol') or ''),'module':module,'screen':screen_context,'manual':knowledge,
                    'verified_draft':result['message'],'required_confidence':result['confidence'],
                },tools=[])
                safe_generated=redact(generated['message'])
                result.update({'message':safe_generated,'speech_text':safe_generated,
                    'provider':generated['provider'],'model':generated['model'],
                    'provider_response_id':generated.get('response_id')})
            except ProviderUnavailable as exc:
                app.logger.warning('LIAM usa recuperación local por indisponibilidad del proveedor: %s',str(exc))
        audit_lia(ctx,'QUESTION_COMPLETED',module=module,request_id=result['request_id'],metadata={'length':len(question),'provider':result['provider']})
        result['ui']=visual_payload(result,module)
        result.update(result['ui'])
        save_exchange(ctx,question=question,answer=result['message'],module=module,request_id=result['request_id'])
        return jsonify(result), 200

    @bp.post('/actions/client-event')
    def client_action_event():
        ctx=get_request_user_context();data=request.get_json(silent=True) or {}
        action=str(data.get('action') or '')
        if action not in {'open_module','search_beneficiary','download_rpp','download_bienestarina','download_ram','generate_monthly_reports','publish_master_database','consolidate_master_database','create_user','update_user','create_foundation','update_foundation'}: return jsonify({'error':'Acción de cliente no registrada.'}),422
        status=str(data.get('status') or '')
        if status not in {'completed','failed','cancelled'}: return jsonify({'error':'Estado de acción no válido.'}),422
        audit_lia(ctx,'CLIENT_ACTION_'+status.upper(),module=str(data.get('module') or '')[:80],tool=action,success=status=='completed',request_id=str(data.get('request_id') or '')[:64],metadata={'detail':redact(str(data.get('detail') or ''))[:160]})
        return jsonify({'ok':True}),200

    @bp.get('/actions/policy')
    def actions_policy():
        ctx=get_request_user_context()
        return jsonify({'role':str(ctx.get('rol') or ''),'actions':public_policy(str(ctx.get('rol') or ''))}),200

    @bp.post('/actions/confirm/<string:proposal_id>')
    def confirm_server_action(proposal_id):
        ctx=get_request_user_context()
        try:result=confirm_credit_proposal(database_path,proposal_id,int(ctx.get('usuario_id') or 0),str(ctx.get('rol') or ''))
        except PermissionError as exc:return jsonify({'error':str(exc)}),403
        except LookupError as exc:return jsonify({'error':str(exc)}),404
        except ValueError as exc:return jsonify({'error':str(exc)}),422
        audit_lia(ctx,'CREDIT_ACTION_COMPLETED',module='facturacion',tool='credit_subscription_update',request_id=proposal_id,metadata={'confirmed':True})
        return jsonify(result),200

    @bp.post('/voice/transcribe')
    def voice_transcribe():
        flags = public_flags()
        if not flags['enabled'] or not flags['voice_enabled'] or not local_speech_enabled():
            return jsonify({'error':'El reconocimiento local de voz no está disponible.'}),503
        ctx = get_request_user_context()
        if limited(ctx): return jsonify({'error':'Demasiadas solicitudes de voz. Espera un momento.'}),429
        uploaded = request.files.get('audio')
        if not uploaded: return jsonify({'error':'No se recibió audio para transcribir.'}),400
        payload = uploaded.read(1_100_000)
        if not payload or len(payload) > 1_000_000: return jsonify({'error':'El fragmento de voz supera el límite permitido.'}),413
        if payload[:4] != b'RIFF' or payload[8:12] != b'WAVE': return jsonify({'error':'El formato de audio no es WAV válido.'}),415
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(prefix='liam_voice_',suffix='.wav',delete=False) as output:
                output.write(payload);temporary=output.name
            text=transcribe_wav(temporary)
            audit_lia(ctx,'VOICE_TRANSCRIBED_LOCAL',module=str(request.form.get('module') or 'dashboard'),metadata={'bytes':len(payload),'characters':len(text)})
            return jsonify({'text':text,'provider':'vosk-local','language':'es'}),200
        except (ValueError,RuntimeError) as exc:
            return jsonify({'error':str(exc)}),422
        finally:
            if temporary:
                try: os.unlink(temporary)
                except OSError: pass

    @bp.post('/voice/realtime/call')
    def realtime_voice_call():
        flags=public_flags();liam_flags=public_liam_flags()
        if not flags['enabled'] or not flags['voice_enabled'] or not liam_flags['realtime_voice_enabled']:
            return jsonify({'error':'La conversación de voz en tiempo real está desactivada.'}),503
        ctx=get_request_user_context()
        if limited(ctx):return jsonify({'error':'Demasiadas solicitudes de voz. Espera un momento.'}),429
        api_key=os.getenv('OPENAI_API_KEY','').strip();model=(os.getenv('LIAM_REALTIME_MODEL') or os.getenv('LIA_REALTIME_MODEL') or '').strip()
        if not api_key or not model:return jsonify({'error':'El proveedor de voz en tiempo real no está configurado.'}),503
        sdp=request.get_data(cache=False,as_text=True)
        if not sdp or len(sdp)>64_000 or not sdp.startswith('v=0'):
            return jsonify({'error':'La oferta WebRTC no es válida.'}),400
        module=str(request.headers.get('X-Liam-Module') or 'dashboard').strip()[:80]
        allowed=set(ROLE_MENU_PERMISSIONS.get(str(ctx.get('rol') or ''),[]))
        if allowed and module not in allowed:module='dashboard'
        manual=manual_for_role(str(ctx.get('rol') or ''),module_id=module)
        safe_context=json.dumps({'rol':ctx.get('rol'),'modulo':module,'manual_operativo':manual},ensure_ascii=False,default=str)[:18_000]
        realtime_tools=[
            {'type':'function','name':'propose_platform_action','description':'Prepara, sin ejecutar, una acción solicitada por voz: RAM, RPP, consolidar/publicar Base Maestra, crear/suspender/reactivar usuarios o fundaciones. Siempre muestra confirmación en la interfaz.','parameters':{'type':'object','properties':{'command':{'type':'string','description':'Orden completa pronunciada por el usuario.'}},'required':['command'],'additionalProperties':False}},
            {'type':'function','name':'get_pending_activities_summary','description':'Consulta actividades pendientes autorizadas por periodo y alcance personal o de equipo.','parameters':{'type':'object','properties':{'period':{'type':'string','description':'Periodo opcional en formato AAAA-MM.'},'scope':{'type':'string','enum':['self','team'],'description':'self para pendientes propios; team para equipo autorizado.'}},'additionalProperties':False}},
            {'type':'function','name':'get_foundation_data_summary','description':'Consulta el panorama institucional consolidado de la fundación activa: perfiles, coordinadores, equipos de talento humano, beneficiarios, grupos etarios, UDS, cargas vigentes por fuente, movimientos y campos incompletos. Nunca consulta otra fundación.','parameters':{'type':'object','properties':{},'additionalProperties':False}},
            {'type':'function','name':'get_monthly_relation_summary','description':'Consulta y visualiza la Relación del Mes de la fundación activa con el formato oficial: UDS, docentes, grupos etarios, usuarios, huevos, cubetas, panales, verduras, olla comunitaria y Bienestarina. Úsala siempre que pidan explicar o mostrar la Relación del Mes.','parameters':{'type':'object','properties':{'period':{'type':'string','description':'Periodo opcional AAAA-MM.'}},'additionalProperties':False}},
            {'type':'function','name':'list_foundation_profiles','description':'Lista perfiles de usuario de la fundación de la sesión activa. Puede filtrar por rol y paginar; nunca consulta otra fundación.','parameters':{'type':'object','properties':{'role':{'type':'string'},'limit':{'type':'integer','minimum':1,'maximum':100},'offset':{'type':'integer','minimum':0}},'additionalProperties':False}},
            {'type':'function','name':'search_foundation_beneficiaries','description':'Busca y filtra beneficiarios de la fundación de la sesión activa por nombre, documento, UDS, grupo etario o estado. Es de solo lectura y nunca cruza fundaciones.','parameters':{'type':'object','properties':{'query':{'type':'string'},'unit':{'type':'string'},'age_group':{'type':'string'},'status':{'type':'string'},'limit':{'type':'integer','minimum':1,'maximum':50},'offset':{'type':'integer','minimum':0}},'additionalProperties':False}},
            {'type':'function','name':'universal_search','description':'Busca por nombre, documento, ID, unidad, docente, usuario, archivo, periodo o estado solamente en recursos autorizados de la fundación activa.','parameters':{'type':'object','properties':{'query':{'type':'string'},'resource':{'type':'string','enum':['all','beneficiaries','profiles','talent','units','documents']},'limit':{'type':'integer','minimum':1,'maximum':50},'offset':{'type':'integer','minimum':0}},'required':['query'],'additionalProperties':False}},
            {'type':'function','name':'get_platform_module_summary','description':'Consulta un resumen operativo autorizado de los módulos institucionales conectados.','parameters':{'type':'object','properties':{'module':{'type':'string','enum':sorted(MODULE_DATASETS)}},'required':['module'],'additionalProperties':False}},
            {'type':'function','name':'get_monthly_health_indicators','description':'Consulta indicadores mensuales de carné de salud, crecimiento y desarrollo, control prenatal, registro civil, perímetro braquial y devuelve anexo nominal autorizado de sobrepeso, desnutrición y riesgo.','parameters':{'type':'object','properties':{'unit':{'type':'string'},'limit':{'type':'integer','minimum':1,'maximum':200}},'additionalProperties':False}},
            {'type':'function','name':'compare_periods','description':'Compara dos cruces mensuales existentes sin estimar datos faltantes y sin cruzar fundaciones.','parameters':{'type':'object','properties':{'period_a':{'type':'string'},'period_b':{'type':'string'}},'required':['period_a','period_b'],'additionalProperties':False}},
            {'type':'function','name':'build_custom_report_preview','description':'Construye una vista previa de reporte usando únicamente tipos y campos seguros predefinidos. No ejecuta SQL generado por IA.','parameters':{'type':'object','properties':{'report':{'type':'string','enum':['unit_coverage','age_distribution','coordinator_coverage']},'fields':{'type':'array','items':{'type':'string','enum':['unit','coordinator','teacher','children_count','age_group','units_count']}},'limit':{'type':'integer','minimum':1,'maximum':200}},'required':['report'],'additionalProperties':False}},
            {'type':'function','name':'supervise_deliverables','description':'Compara entregables esperados y recibidos; clasifica pendientes, vencidos, devueltos, incompletos y duplicados dentro del alcance del usuario.','parameters':{'type':'object','properties':{'period':{'type':'string'}},'additionalProperties':False}},
            {'type':'function','name':'get_system_health','description':'Exclusivo de SUPERADMIN. Consulta un diagnóstico sanitizado de base, auditoría, documentos, calendario y Base Maestra sin exponer secretos.','parameters':{'type':'object','properties':{},'additionalProperties':False}},
            {'type':'function','name':'get_foundation_portfolio','description':'Exclusivo de SUPERADMIN. Consulta fundaciones, usuarios, actividad, vencimientos, planes y créditos con alcance global explícito.','parameters':{'type':'object','properties':{'limit':{'type':'integer','minimum':1,'maximum':200}},'additionalProperties':False}},
            {'type':'function','name':'analyze_master_data_quality','description':'Analiza documentos duplicados, campos faltantes, UDS inexistentes, unidades sin responsable y docentes sin unidad. Nunca corrige automáticamente.','parameters':{'type':'object','properties':{},'additionalProperties':False}},
            {'type':'function','name':'get_early_warnings','description':'Consolida señales de riesgo verificables de entregables y Base Maestra. No presenta predicciones como hechos.','parameters':{'type':'object','properties':{'period':{'type':'string'}},'additionalProperties':False}},
            {'type':'function','name':'get_incident_center','description':'Consulta incidencias sanitizadas. Usuario común ve las propias; gerente y SUPERADMIN ven su fundación.','parameters':{'type':'object','properties':{'incident_id':{'type':'string'},'status':{'type':'string','enum':['','OPEN','IN_ANALYSIS','IN_PROGRESS','RESOLVED','CLOSED']},'limit':{'type':'integer','minimum':1,'maximum':100}},'additionalProperties':False}},
            {'type':'function','name':'get_notification_center','description':'Unifica avisos pendientes de planeación, calendario, incidencias, documentos y créditos dentro del alcance autenticado.','parameters':{'type':'object','properties':{'limit':{'type':'integer','minimum':1,'maximum':100}},'additionalProperties':False}},
            {'type':'function','name':'get_structured_error','description':'Explica un código de error de la plataforma.','parameters':{'type':'object','properties':{'code':{'type':'string'}},'required':['code'],'additionalProperties':False}},
            {'type':'function','name':'get_document_processing_status','description':'Consulta el estado autorizado de un documento procesado.','parameters':{'type':'object','properties':{'document_id':{'type':'integer'}},'required':['document_id'],'additionalProperties':False}},
            {'type':'function','name':'get_format_generation_status','description':'Consulta el estado de una generación de formato.','parameters':{'type':'object','properties':{'test_id':{'type':'integer'}},'required':['test_id'],'additionalProperties':False}},
        ]
        voice_policy='; '.join(f"{item['action']}={item['risk']}/{item['confirmation']}" for item in public_policy(str(ctx.get('rol') or '')) if item.get('allowed') and item.get('connected'))
        session={'type':'realtime','model':model,'instructions':realtime_instructions(
            action_policy=voice_policy, authorized_context=safe_context
        ),'output_modalities':['audio'],'tools':realtime_tools,'tool_choice':'auto','audio':{
            'input':{'transcription':{'model':'gpt-4o-mini-transcribe','language':'es'},'turn_detection':{'type':'server_vad','create_response':True,'interrupt_response':True}},
            'output':{'voice':'marin'},
        }}
        safety_id=hashlib.sha256(f"liam:{ctx.get('fundacion_id')}:{ctx.get('usuario_id')}".encode()).hexdigest()
        try:
            upstream=requests.post('https://api.openai.com/v1/realtime/calls',headers={'Authorization':f'Bearer {api_key}','OpenAI-Safety-Identifier':safety_id},files={'sdp':(None,sdp),'session':(None,json.dumps(session,ensure_ascii=False))},timeout=(8,30))
        except requests.RequestException as exc:
            app.logger.warning('LIAN Realtime no pudo iniciar: %s',type(exc).__name__)
            return jsonify({'error':'No fue posible conectar la conversación de voz con el proveedor.'}),503
        if not upstream.ok:
            try:detail=(upstream.json().get('error') or {}).get('message')
            except ValueError:detail=None
            app.logger.warning('LIAN Realtime rechazado: status=%s',upstream.status_code)
            return jsonify({'error':detail or 'El proveedor rechazó la sesión de voz.'}),upstream.status_code
        audit_lia(ctx,'REALTIME_VOICE_STARTED',module=module,metadata={'model':model})
        return Response(upstream.text,200,{'Content-Type':'application/sdp','Cache-Control':'no-store'})

    @bp.post('/voice/realtime/event')
    def realtime_voice_event():
        ctx=get_request_user_context();data=request.get_json(silent=True) or {};event=str(data.get('event') or '')
        if event=='transcript':
            role=str(data.get('role') or '');content=str(data.get('content') or '').strip();module=str(data.get('module') or 'dashboard')[:80]
            if role not in {'user','assistant'} or not content:return jsonify({'error':'Transcripción de voz no válida.'}),422
            save_message(ctx,role=role,content=content,module=module,request_id=str(data.get('request_id') or uuid.uuid4().hex)[:64])
            return jsonify({'ok':True,'saved':True}),201
        if event not in {'ended','failed'}:return jsonify({'error':'Evento de voz no válido.'}),422
        duration=max(0,min(600,int(data.get('duration') or 0)));reason=str(data.get('reason') or '')[:60]
        audit_lia(ctx,'REALTIME_VOICE_'+event.upper(),module=str(data.get('module') or 'dashboard')[:80],success=event=='ended',metadata={'duration_seconds':duration,'reason':reason})
        return jsonify({'ok':True}),200

    @bp.get('/health')
    def health():
        flags = public_flags()
        provider=provider_status()
        local_voice=local_speech_status()
        mode='provider' if provider['ready'] else 'institutional_static'
        return jsonify({'status':'ok','enabled':flags['enabled'],'mode':mode,'provider_ready':provider['ready'],
            'components':{'text':flags['text_enabled'],'context':flags['context_help_enabled'],
                'tours':flags['guided_tours_enabled'],'voice':flags['voice_enabled'],
                'local_voice':local_voice,
                'generative_ai':flags['ai_enabled'] and provider['ready'],
                'realtime_voice':public_liam_flags()['realtime_voice_enabled'] and bool((os.getenv('LIAM_REALTIME_MODEL') or os.getenv('LIA_REALTIME_MODEL') or '').strip())},
            'operational':flags['enabled'] and flags['text_enabled'],
            'degraded_reasons':([] if flags['enabled'] else ['assistant_disabled'])+
                ([] if flags['text_enabled'] else ['text_disabled'])+
                ([] if not flags['ai_enabled'] or provider['ready'] else ['provider_not_ready'])}), 200

    @bp.get('/errors/<string:incident_id>')
    def error_diagnosis(incident_id):
        ctx=get_request_user_context();item=get_incident(database_path,incident_id,int(ctx.get('fundacion_id') or 1))
        if not item:return jsonify({'error':'Incidente no encontrado o no autorizado.'}),404
        return jsonify({'incident':item}),200

    @bp.get('/errors')
    def error_center_list():
        ctx=get_request_user_context()
        if str(ctx.get('rol') or '') not in {'SUPERADMIN','GERENTE','COORDINADOR'}:return jsonify({'error':'No tienes permiso para consultar el centro de diagnóstico.'}),403
        return jsonify({'incidents':list_incidents(database_path,int(ctx.get('fundacion_id') or 1),request.args.get('limit',50,type=int))}),200

    @bp.get('/tools')
    def tools_available():
        if not public_flags()['enabled']: return jsonify({'error':'LÍA está desactivada.'}),404
        get_request_user_context()
        return jsonify({'tools':sorted(ALLOWED_TOOLS),'write_tools':[],'proposal_tools':['propose_platform_action']}),200

    @bp.post('/tools/<string:tool_name>')
    def run_tool(tool_name: str):
        if not public_flags()['enabled']: return jsonify({'error':'LÍA está desactivada.'}),404
        ctx=get_request_user_context(); user=dict(getattr(g,'current_user',None) or {})
        if limited(ctx): return jsonify({'error':'Demasiadas solicitudes a LÍA. Espera un momento.'}),429
        if not user.get('id'): user={'id':ctx.get('usuario_id'),'rol':ctx.get('rol')}
        request_id=uuid.uuid4().hex
        try:
            outcome=orchestrator.run(tool_name,args=request.get_json(silent=True) or {},tenant_id=int(ctx.get('fundacion_id') or 1),user=user,module=str(request.headers.get('X-Liam-Module') or 'dashboard'),request_id=request_id);result=outcome.result
        except PermissionError as exc: audit_lia(ctx,'TOOL_REJECTED',tool=tool_name,success=False,request_id=request_id);return jsonify({'error':str(exc),'request_id':request_id}),403
        except LookupError as exc: audit_lia(ctx,'TOOL_NOT_FOUND',tool=tool_name,success=False,request_id=request_id);return jsonify({'error':str(exc),'request_id':request_id}),404
        except ValueError as exc: audit_lia(ctx,'TOOL_INVALID_ARGUMENT',tool=tool_name,success=False,request_id=request_id);return jsonify({'error':str(exc),'request_id':request_id}),422
        audit_lia(ctx,'TOOL_COMPLETED',tool=tool_name,request_id=request_id,metadata={**outcome.telemetry,'proposal_only':tool_name=='propose_platform_action'})
        active_module=str(request.headers.get('X-Liam-Module') or 'dashboard').strip()[:80]
        ui=visual_payload({'message':'Datos consultados por Lía.','tool_result':result},active_module)
        return jsonify({'tool':tool_name,'result':result,'ui':ui,'read_only':outcome.telemetry['read_only'],'request_id':request_id,'trace':outcome.telemetry}),200

    @bp.route('/preferences',methods=['GET','PUT'])
    def preferences():
        if not public_flags()['enabled']: return jsonify({'error':'LÍA está desactivada.'}),404
        ctx=get_request_user_context();fid=int(ctx.get('fundacion_id') or 1);uid=int(ctx.get('usuario_id') or 0);conn=connect()
        if request.method=='GET':
            row=conn.execute('SELECT voice_enabled,auto_speak_enabled,muted,speech_rate,reduced_motion,language FROM lia_user_preferences WHERE fundacion_id=? AND usuario_id=?',(fid,uid)).fetchone();conn.close()
            return jsonify({'preferences':dict(row) if row else {'voice_enabled':0,'auto_speak_enabled':0,'muted':0,'speech_rate':.95,'reduced_motion':0,'language':'es-CO'}}),200
        data=request.get_json(silent=True) or {}
        try: rate=max(.6,min(1.5,float(data.get('speech_rate') or .95)))
        except (TypeError,ValueError): conn.close();return jsonify({'error':'La velocidad de voz no es válida.'}),422
        now=datetime.now().isoformat(timespec='seconds')
        values=(1 if data.get('voice_enabled') else 0,1 if data.get('auto_speak_enabled') else 0,1 if data.get('muted') else 0,rate,1 if data.get('reduced_motion') else 0,'es-CO')
        conn.execute('''INSERT INTO lia_user_preferences(fundacion_id,usuario_id,voice_enabled,auto_speak_enabled,muted,speech_rate,reduced_motion,language,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(fundacion_id,usuario_id) DO UPDATE SET voice_enabled=excluded.voice_enabled,auto_speak_enabled=excluded.auto_speak_enabled,muted=excluded.muted,speech_rate=excluded.speech_rate,reduced_motion=excluded.reduced_motion,language=excluded.language,updated_at=excluded.updated_at''',(fid,uid,*values,now,now));conn.commit();conn.close();audit_lia(ctx,'PREFERENCES_UPDATED',metadata={'voice_enabled':bool(values[0]),'muted':bool(values[2])})
        return jsonify({'message':'Preferencias de LÍA actualizadas.'}),200

    @bp.post('/feedback')
    def feedback():
        flags=public_flags()
        if not flags['enabled'] or not flags['feedback_enabled']: return jsonify({'error':'La retroalimentación está desactivada.'}),404
        ctx=get_request_user_context();data=request.get_json(silent=True) or {};rating=int(data.get('rating') or 0)
        if rating not in {-1,1}: return jsonify({'error':'Valoración no válida.'}),422
        reason=str(data.get('reason') or '')[:240];module=str(data.get('module') or '')[:80];request_id=str(data.get('request_id') or '')[:64];conn=connect();now=datetime.now().isoformat(timespec='seconds')
        conn.execute('INSERT INTO lia_feedback(fundacion_id,usuario_id,request_id,rating,reason,module,created_at) VALUES(?,?,?,?,?,?,?)',(int(ctx.get('fundacion_id') or 1),int(ctx.get('usuario_id') or 0),request_id,rating,reason,module,now));conn.commit();conn.close();audit_lia(ctx,'FEEDBACK_RECORDED',module=module,request_id=request_id,metadata={'rating':rating})
        return jsonify({'message':'Gracias. Registramos tu valoración sin guardar datos personales de la conversación.'}),201

    app.register_blueprint(bp)

    @app.after_request
    def liam_error_center(response):
        try:
            user=dict(getattr(g,'current_user',None) or {})
            if not user or not request.path.startswith('/api/') or request.path.startswith('/api/asistente-capacitacion/errors') or response.status_code<400 or response.status_code in {401}:
                return response
            payload=response.get_json(silent=True) if response.is_json else None
            if not isinstance(payload,dict):return response
            message=str(payload.get('error') or payload.get('mensaje') or payload.get('message') or '')
            code=str(payload.get('code') or payload.get('codigo') or f'HTTP_{response.status_code}')
            module=request.path.split('/')[2] if len(request.path.split('/'))>2 else 'plataforma'
            incident=record_incident(database_path,tenant_id=user.get('fundacion_id'),user_id=user.get('id'),module=module,action=request.endpoint or request.path,status=response.status_code,code=code,message=message,request_id=payload.get('request_id') or payload.get('trace_id'),context={'method':request.method})
            payload['incident_id']=incident['incident_id'];payload['diagnostic']={k:incident[k] for k in ('type','cause','solution','severity','safe_retry','auto_correctable')}
            response.set_data(json.dumps(payload,ensure_ascii=False));response.headers['Content-Type']='application/json; charset=utf-8';response.headers['X-Liam-Incident-ID']=incident['incident_id']
        except Exception as exc:
            app.logger.warning('Centro de errores LIAM continuó sin registrar: %s',type(exc).__name__)
        return response
