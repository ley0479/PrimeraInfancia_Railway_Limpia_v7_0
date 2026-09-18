from pathlib import Path
import json
import sys

from docx import Document
from openpyxl import load_workbook


def clean(value):
    return " ".join(str(value or "").replace("\n", " | ").split())


def docx_info(path):
    doc = Document(path)
    paragraphs = [clean(p.text) for p in doc.paragraphs if clean(p.text)]
    tables = []
    for index, table in enumerate(doc.tables, 1):
        rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
        nonempty = [row for row in rows if any(row)]
        tables.append({
            "index": index,
            "rows": len(rows),
            "cols": max((len(r) for r in rows), default=0),
            "sample": nonempty[:12],
        })
    images = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)
    sections = [{
        "orientation": str(section.orientation),
        "width_cm": round(section.page_width.cm, 2),
        "height_cm": round(section.page_height.cm, 2),
        "margins_cm": [round(section.top_margin.cm, 2), round(section.right_margin.cm, 2), round(section.bottom_margin.cm, 2), round(section.left_margin.cm, 2)],
    } for section in doc.sections]
    return {"type": "docx", "paragraphs": paragraphs[:80], "tables": tables, "images": images, "sections": sections}


def xlsx_info(path):
    wb = load_workbook(path, read_only=False, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        populated = []
        for row in ws.iter_rows():
            values = [clean(cell.value) for cell in row]
            if any(values):
                populated.append({"row": row[0].row, "values": values[:25]})
            if len(populated) >= 18:
                break
        sheets.append({
            "name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column,
            "merged": [str(r) for r in list(ws.merged_cells.ranges)[:30]], "sample": populated,
        })
    return {"type": "xlsx", "sheets": sheets}


root = Path(sys.argv[1])
compact = "--compact" in sys.argv
for path in sorted(root.iterdir(), key=lambda p: p.name.casefold()):
    if not path.is_file() or path.suffix.lower() not in {".docx", ".xlsx"}:
        continue
    try:
        info = docx_info(path) if path.suffix.lower() == ".docx" else xlsx_info(path)
        if compact and info["type"] == "docx":
            corpus = " ".join(info.get("paragraphs", [])) + " " + " ".join(
                " ".join(" ".join(row) for row in table.get("sample", [])) for table in info.get("tables", [])
            )
            wanted = ["ACTA DE REUNIONES", "FECHA", "LUGAR", "RESPONSABLE", "TIPO DE ACTIVIDAD", "DIRIGIDO A", "TEMA", "OBJETIVO", "AGENDA", "DESARROLLO DE LA ACTIVIDAD", "COMPROMISOS", "REGISTROS FOTOGRAFICOS", "ASISTENCIA", "NOVEDAD", "FIRMAS"]
            info = {"type": "docx", "paragraph_count": len(info.get("paragraphs", [])), "tables": [f'{x["rows"]}x{x["cols"]}' for x in info.get("tables", [])], "images": info.get("images", 0), "fields": [x for x in wanted if x in corpus.upper()]}
        elif compact:
            info = {"type": "xlsx", "sheets": len(info.get("sheets", [])), "sheet_names": [x["name"] for x in info.get("sheets", [])], "dimensions": [f'{x["max_row"]}x{x["max_column"]}' for x in info.get("sheets", [])], "headers": [x.get("sample", [])[:4] for x in info.get("sheets", [])[:2]]}
        print(json.dumps({"file": path.name, **info}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"file": path.name, "error": str(exc)}, ensure_ascii=False))
